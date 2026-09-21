"""Reusable application export; snapshots, not recalculated business records."""
import io
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .store import CHINA


LABELS = {"pending": "待处理", "waiting_merchant_reply": "等待商家回复", "success": "处理成功", "exception": "异常", "skipped": "已跳过（未完全匹配）", "blocked": "屏蔽", "risk": "风险"}


def numeric(value):
    if value in (None, "") or isinstance(value, bool):
        return None
    try:
        number = Decimal(str(value))
        return float(number) if number.is_finite() else None
    except InvalidOperation:
        return None


def execution_xlsx(rows, batch=None):
    batch = batch or {}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品执行情况"
    sheet.sheet_view.showGridLines = False
    headers = ["ERP商品ID", "商品标题", "执行结果", "跳过、失败或暂停原因", "目标SKU", "供货成本（人民币）",
               "包装重量（克）", "净收益（美元）", "修改前重量（克）", "修改前净收益（美元）",
               "计划修改重量（克）", "计划修改净收益（美元）", "修改后回读重量（克）", "修改后回读净收益（美元）",
               "保存回读验证", "成本来源", "重量来源", "美元汇率（人民币/美元）", "汇率日期", "1688货源链接", "记录时间（北京时间）",
               "修改前智赢状态", "计划修改智赢状态", "修改后智赢状态",
               "当前ERP重量（克）", "当前ERP净收益（美元）", "核验方式", "人工核验备注"]
    sheet["A2"] = "AI核重核价 · 产品执行情况"
    sheet["A2"].font = Font(name="Microsoft YaHei", size=14, bold=True, color="17365D")
    sheet.row_dimensions[2].height = 30
    successes = sum(row.get("execution_result", LABELS.get(row["status"])) == "处理成功" for row in rows)
    skipped = sum(row["status"] == "skipped" for row in rows)
    blocked, risk = (sum(row['status'] == status for row in rows) for status in ('blocked', 'risk'))
    sheet["A3"] = f"共 {len(rows)} 件，成功 {successes} 件，屏蔽 {blocked} 件，风险 {risk} 件，跳过 {skipped} 件；批次：{batch.get('run_id', '全部记录')}"
    sheet["A4"] = "记录为执行时快照。修改后数据仅来自保存后的页面回读；空白表示尚未读取或未执行，不能视为0。"
    sheet.row_dimensions[4].height = 26
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(5, column, label)
        cell.fill = PatternFill("solid", fgColor="24476C")
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    sheet.row_dimensions[5].height = 42
    numeric_columns = set(range(6, 15)) | {18, 25, 26}
    for index, row in enumerate(rows, 6):
        before, intent, after = (row.get(field) or {} for field in ("erp_before", "write_intent", "erp_after"))
        source, pricing = row.get("info_sources") or {}, row.get("pricing") or {}
        stamp = row.get("saved_at") or row.get("updated_at")
        failure = "：".join(str(row.get(field) or "") for field in ("exception_reason", "exception_detail") if row.get(field)) if row["status"] == "exception" else ""
        current = after or before or {"weight_g": row.get("reference_weight_g")}
        values = [str(row["erp_goods_id"]), row.get("title"), row.get("execution_result") or LABELS[row["status"]],
                  failure or row.get("execution_reason") or row.get("decision_reason") or row.get("skip_reason") or row.get("defer_reason"),
                  row.get("erp_sku"), row.get("cost_price"), row.get("weight_g"), row.get("net_income_usd"),
                  before.get("weight_g"), before.get("net_income_usd"), intent.get("weight_g"), intent.get("net_income_usd"),
                  after.get("weight_g"), after.get("net_income_usd"),
                  "通过" if row.get("write_verified") is True else "未确认", source.get("cost_price"), source.get("weight_g"),
                  pricing.get("cny_per_usd"), pricing.get("rate_date"), row.get("supplier_url"),
                  datetime.fromtimestamp(stamp, CHINA).replace(tzinfo=None) if stamp else None,
                  before.get("review_status"), intent.get("review_status"), after.get("review_status"),
                  current.get("weight_g"), current.get("net_income_usd"),
                  "人工核验" if row.get("verification_mode") == "manual" else "自动核验",
                  (row.get("manual_verification") or {}).get("note")]
        for column, value in enumerate(values, 1):
            if column in numeric_columns:
                value = numeric(value)
            elif isinstance(value, str):
                value = re.sub(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]", "", value)[:32767]
            cell = sheet.cell(index, column, value)
            # Text is untrusted product/reply content, never an Excel formula.
            if isinstance(value, str):
                cell.data_type = "s"
            cell.font = Font(name="Microsoft YaHei", size=10, color="243746")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F5FA")
            if column in numeric_columns:
                cell.number_format = "0.000000" if column == 18 else "0.##"
            if column == 21:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
        sheet.row_dimensions[index].height = 64
        sheet.cell(index, 3).font = Font(name="Microsoft YaHei", bold=True, color="237344" if values[2] == "处理成功" else "9C3B24")
    widths = [20, 48, 18, 68, 32] + [18] * 10 + [18, 18, 24, 18, 52, 24, 20, 20, 20, 18, 18, 20, 40]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "C6"
    sheet.auto_filter.ref = f"A5:AB{max(5, len(rows) + 5)}"
    sheet.print_title_rows = "1:5"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
