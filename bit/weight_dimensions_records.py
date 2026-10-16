"""Uploaded order measurements and explicit Mercado Libre listing updates."""
from __future__ import annotations

import io
import json
import logging
import os
import re
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from erp.package_update_weight import package_update_weight_g

from bit import bit_mysql
from bit import bit_db_api
from bit.shipment_compensation import apply_dimension_costs
from bit.bit_store_link_sync import _client_and_token
from mercado_api.client import MercadoLibreClient


MAX_UPLOAD_BYTES = 16 * 1024 * 1024
MAX_ORDER_ROWS = 3000
_lock = threading.RLock()
_tasks: dict[str, dict[str, Any]] = {}
_latest_task_by_owner: dict[str, str] = {}
_execution_threads: dict[str, threading.Thread] = {}
_executions_shutting_down = False
_logger = logging.getLogger(__name__)
READ_MAX_WORKERS = 8
UPDATE_MAX_WORKERS = 4


class _ExecutionStopped(Exception):
    """Cancel discovery before issuing the next remote request."""


HEADERS = {
    "time": ("时间", "下单时间"),
    "order_number": ("编号", "订单号", "销售编号"),
    "salesperson": ("业务员",),
    "source": ("来源",),
    "company_store": ("公司店铺", "店铺"),
    "product_id": ("产品id", "产品ID", "产品号"),
    "category": ("产品分类", "产品类别"),
    "title": ("标题", "产品标题"),
    "image_url": ("图片", "主图"),
    "shipment_id": ("运单号",),
    "tracking_number": ("追踪号", "物流单号"),
    "carrier": ("运输商", "物流商"),
    "region": ("地区", "站点"),
}

PUBLIC_COLUMNS = (
    ("time", "时间"),
    ("freight_changed_at", "运费变更时间"),
    ("image_url", "图片"),
    ("order_number", "编号"),
    ("salesperson", "业务员"),
    ("source", "来源"),
    ("company_store", "公司店铺"),
    ("product_id", "产品id"),
    ("category", "产品分类"),
    ("title", "标题"),
    ("shipment_id", "运单号"),
    ("tracking_number", "追踪号"),
    ("carrier", "运输商"),
    ("region", "地区"),
    ("declared_weight_g", "当前标记重量（克）"),
    ("declared_dimensions_cm", "当前标记尺寸（厘米）"),
    ("declared_freight", "当前运费"),
    ("actual_weight_g", "实际重量（克）"),
    ("actual_dimensions_cm", "实际尺寸（厘米）"),
    ("actual_freight", "实际运费"),
    ("freight_difference", "运费差值（实际-当前）"),
    ("marketplace_item_id", "美客多产品编号"),
    ("query_status", "查询状态"),
    ("query_error", "说明"),
)
EXPORT_COLUMNS = PUBLIC_COLUMNS + (
    ("current_net_proceeds_usd", "本次净收益（USD）"),
    ("_zeshun_status", "泽顺 ERP 执行状态"),
    ("_zying_status", "智赢产品执行状态"),
    ("execution_error", "执行说明"),
    ("execution_logs", "执行日志"),
)

_COUNTRY_TO_SITE = {
    "墨西哥": "MLM", "巴西": "MLB", "智利": "MLC", "哥伦比亚": "MCO",
    "阿根廷": "MLA", "乌拉圭": "MLU", "秘鲁": "MPE", "厄瓜多尔": "MEC",
}
_PACKAGE_ATTRIBUTES = {
    "PACKAGE_WEIGHT": ("weight", "g"),
    "PACKAGE_LENGTH": ("length", "cm"),
    "PACKAGE_WIDTH": ("width", "cm"),
    "PACKAGE_HEIGHT": ("height", "cm"),
}
_SITE_NAME_BY_ID = {
    **{site_id: site_name for site_id, site_name in getattr(
        bit_mysql, "MERCADO_CONFIGURABLE_SITES", {}
    ).items()},
    "MLM": "墨西哥", "MLB": "巴西", "MLC": "智利", "MCO": "哥伦比亚",
    "MLA": "阿根廷", "MLU": "乌拉圭", "MPE": "秘鲁", "MEC": "厄瓜多尔",
}


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def _normalize_order_time(value: str) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    candidate = raw.replace("/", "-").replace("年", "-").replace("月", "-").replace("日", " ")
    try:
        return datetime.fromisoformat(candidate).isoformat(sep=" ", timespec="seconds")
    except ValueError:
        for pattern in ("%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y%m%d%H%M%S", "%Y%m%d"):
            try:
                return datetime.strptime(candidate.strip(), pattern).isoformat(sep=" ", timespec="seconds")
            except ValueError:
                pass
    return raw


def read_order_file(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise ValueError("文件不能超过 16 MB")
    if not str(filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("请上传 .xlsx 或 .xlsm 订单文件")
    book = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        rows = book.active.values
        header = next(rows, None)
        if not header:
            raise ValueError("订单文件没有表头")
        header_map = {str(value or "").strip().casefold(): index for index, value in enumerate(header)}
        columns: dict[str, int] = {}
        missing = []
        for field, aliases in HEADERS.items():
            index = next((header_map.get(name.casefold()) for name in aliases if name.casefold() in header_map), None)
            if index is None and field in {"order_number", "company_store"}:
                missing.append(aliases[0])
            if index is not None:
                columns[field] = index
        if missing:
            raise ValueError("订单文件缺少必需列：" + "、".join(missing))
        records = []
        for row_number, row in enumerate(rows, start=2):
            record = {field: _cell_text(row[index] if index < len(row) else None)
                      for field, index in columns.items()}
            record["order_number"] = re.sub(r"\.0$", "", record.get("order_number", ""))
            record["time"] = _normalize_order_time(record.get("time", ""))
            if not record["order_number"]:
                continue
            if len(records) >= MAX_ORDER_ROWS:
                raise ValueError(f"每次最多处理 {MAX_ORDER_ROWS} 条订单")
            record.update({
                "source_row": row_number,
                "declared_weight_g": "",
                "declared_dimensions_cm": "",
                "declared_freight": "",
                "actual_weight_g": "",
                "actual_dimensions_cm": "",
                "actual_freight": "",
                "marketplace_item_id": "",
                "query_status": "等待查询",
                "query_error": "",
                "_token_id": None,
                "_global_item_id": "",
                "_net_proceeds_usd": None,
                "_package_status": "",
                "_net_status": "",
            })
            records.append(record)
        if not records:
            raise ValueError("文件中没有带订单编号的有效订单行")
        return records
    finally:
        book.close()


def _store_map(store_rows: list[Mapping[str, Any]], allowed_token_ids: set[int] | None):
    by_alias: dict[str, list[Mapping[str, Any]]] = defaultdict(list)
    for store in store_rows:
        token_id = int(store.get("id") or 0)
        if not token_id or not store.get("enabled", True):
            continue
        if allowed_token_ids is not None and token_id not in allowed_token_ids:
            continue
        for alias in (store.get("display_name"), store.get("nickname")):
            key = str(alias or "").strip().casefold()
            if key and store not in by_alias[key]:
                by_alias[key].append(store)
    return by_alias


def _resolve_store(record: Mapping[str, Any], store_map):
    full_name = str(record.get("company_store") or "").strip()
    suffix = re.search(r"\(([^()]*)\)\s*$", full_name)
    country = suffix.group(1).strip() if suffix else ""
    base = re.sub(r"\s*\([^()]*\)\s*$", "", full_name).strip()
    candidates = list(store_map.get(base.casefold()) or store_map.get(full_name.casefold()) or [])
    if not candidates:
        return None
    site_id = _COUNTRY_TO_SITE.get(country)
    if site_id:
        site_candidates = [
            row for row in candidates
            if site_id in {str(setting.get("site_id") or "").upper()
                           for setting in row.get("site_settings") or []}
        ]
        if site_candidates:
            candidates = site_candidates
    return candidates[0] if len(candidates) == 1 else None


def _measurement(data: Mapping[str, Any], package_key: str) -> tuple[str, str]:
    """Keep valid weight and dimensions independently; never substitute billable weight."""
    package = (data.get("package") or {}).get(package_key) or {}
    weight = package.get("weight") or package.get("weights") or {}
    dims = package.get("dimensions") or {}

    def positive(raw):
        try:
            value = Decimal(str(raw))
        except (InvalidOperation, TypeError, ValueError):
            return ""
        return format(value.normalize(), "f") if value.is_finite() and value > 0 else ""

    net = positive(weight.get("net")) if str(weight.get("unit") or "").lower() == "g" else ""
    sides = [positive(dims.get(key)) for key in ("length", "width", "height")]
    dimensions = "x".join(sides) if all(sides) and str(dims.get("unit") or "").lower() == "cm" else ""
    return net, dimensions


def _money(costs: Mapping[str, Any], section: str) -> str:
    sender = ((costs.get(section) or {}).get("sender") or {})
    value = sender.get("cost")
    if value in (None, ""):
        return ""
    currency = str(costs.get("currency_id") or "").strip()
    return f"{value} {currency}".strip()


def _latest_order_row(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Choose the newest order, with stable tie-breakers when timestamps match."""
    return max(
        rows,
        key=lambda row: (
            _normalize_order_time(str(row.get("time") or "")),
            str(row.get("order_number") or ""),
            int(row.get("source_row") or 0),
        ),
    )


def _match_order_item(
    client: MercadoLibreClient,
    pack: Mapping[str, Any],
    product_id: str,
    marketplace_item_id: str = "",
):
    wanted = str(product_id or "").strip().casefold()
    wanted_marketplace = str(marketplace_item_id or "").strip().casefold()
    candidates = []
    for order_ref in pack.get("orders") or []:
        order_id = str((order_ref or {}).get("id") or "").strip()
        if not order_id:
            continue
        order = client.get_order(order_id)
        for order_line in order.get("order_items") or []:
            item = order_line.get("item") or {}
            candidates.append((item, order_line, order))
    if not candidates:
        raise ValueError("订单明细中没有可更新的产品")
    matching = [
        candidate for candidate in candidates
        if (wanted and wanted in {
            str(candidate[0].get("seller_sku") or "").strip().casefold(),
            str(candidate[0].get("id") or "").strip().casefold(),
            str(candidate[0].get("parent_item_id") or "").strip().casefold(),
        }) or (
            wanted_marketplace
            and wanted_marketplace == str(candidate[0].get("id") or "").strip().casefold()
        )
    ]
    if len(matching) == 1:
        return matching[0]
    if len(candidates) == 1:
        return candidates[0]
    raise ValueError("包裹含多个产品，且无法用产品 id 唯一匹配商品明细")


def _extract_net_proceeds(item: Mapping[str, Any]) -> Decimal | None:
    values = item.get("net_proceeds")
    for candidate in values if isinstance(values, list) else [values]:
        if not isinstance(candidate, Mapping):
            continue
        if str(candidate.get("currency_id") or "").upper() != "USD":
            continue
        try:
            amount = Decimal(str(candidate.get("amount")))
        except (InvalidOperation, TypeError, ValueError):
            continue
        if amount.is_finite() and amount > 0:
            return amount
    return None


def _read_measurements(record, client, shipment_id):
    """Read and retain measurements even when other package fields are absent."""
    record["query_error"] = ""
    notes = {}
    try:
        compensation = client.request(
            "GET", f"/marketplace/shipments/{shipment_id}/compensation_costs",
            params={"weight_unit": "g", "dimensions_unit": "cm"},
        )
    except Exception as exc:
        compensation = {}
        record["query_error"] = f"官方补差数据读取失败：{str(exc)[:200]}"
    record["_compensation_package"] = compensation.get("package") or {}
    for section, prefix in (("declared", "declared"), ("validated", "actual")):
        weight, dimensions = _measurement(compensation, section)
        record[f"{prefix}_weight_g"] = weight
        record[f"{prefix}_dimensions_cm"] = dimensions
        for suffix, value in (("weight_g", weight), ("dimensions_cm", dimensions)):
            if not value:
                notes[f"{prefix}_{suffix}"] = "官方未提供" if compensation else "补差数据不可用"
        freight = _money(compensation.get("costs") or {}, section)
        if freight and not record.get("_dimension_compensation_confirmed"):
            record[f"{prefix}_freight"] = freight
    # Shipment dimensions describe the shipment's declared package. They must
    # never be presented as the carrier-validated actual measurements.
    if not record["declared_weight_g"] or not record["declared_dimensions_cm"]:
        try:
            shipment_dims = (client.get_shipment(shipment_id) or {}).get("dimensions") or {}
            record["_shipment_dimensions"] = shipment_dims
            weight, dimensions = _measurement({"package": {"declared": {
                "weight": {"net": shipment_dims.get("weight"), "unit": "g"},
                "dimensions": {**shipment_dims, "unit": "cm"},
            }}}, "declared")
            for key, value in (("declared_weight_g", weight), ("declared_dimensions_cm", dimensions)):
                if not record[key] and value:
                    record[key] = value
                    notes[key] = "来自官方运单申报数据"
        except Exception as exc:
            record["query_error"] += f"；官方运单尺寸读取失败：{str(exc)[:120]}"
    record["measurement_notes"] = notes
    record["_measurement_version"] = 2
    record["_measurements_checked_at"] = datetime.now().isoformat(timespec="seconds")
    record["shipment_id"] = shipment_id



def _read_one_record(record, store_map, token_clients, token_errors, worker_local):
    """Read package compensation and marketplace identifiers for one order."""
    token_id = int(record.get("_token_id") or 0)
    if not token_id:
        raise ValueError("订单没有可用的店铺授权")
    if token_id in token_errors:
        raise ValueError("店铺授权不可用：" + token_errors[token_id])
    clients = getattr(worker_local, "clients", None)
    if clients is None:
        clients = worker_local.clients = {}
    client = clients.get(token_id)
    if client is None:
        template = token_clients.get(token_id)
        if template is None:
            raise ValueError("店铺授权不可用")
        client = MercadoLibreClient(template.access_token, timeout=template.timeout)
        clients[token_id] = client

    order_number = str(record.get("order_number") or "").strip()
    pack = client.request("GET", f"/marketplace/orders/pack/{order_number}")
    if str(pack.get("id")) != order_number:
        raise ValueError("接口返回的包裹号与订单编号不一致")
    shipment_id = str((pack.get("shipment") or {}).get("id") or "").strip()
    if not shipment_id:
        raise ValueError("订单包裹没有运单号")
    # Uploads have no cached official costs; background candidates already do.
    if "_dimension_compensation_confirmed" not in record:
        apply_dimension_costs(record, client.get_shipment_costs(shipment_id))
    _read_measurements(record, client, shipment_id)

    item = {}
    match_error = ""
    try:
        item, _order_line, _order = _match_order_item(
            client, pack, record.get("product_id", ""), record.get("marketplace_item_id", "")
        )
    except Exception as exc:
        match_error = str(exc)
    marketplace_item_id = str(item.get("id") or record.get("marketplace_item_id") or "").strip()
    global_item_id = str(item.get("parent_item_id") or record.get("_global_item_id") or "").strip()
    item_detail = {}
    if marketplace_item_id:
        try:
            item_detail = client.get_marketplace_item(
                marketplace_item_id,
                attributes=("id", "title", "secure_thumbnail", "thumbnail", "attributes", "net_proceeds", "cbt_item_id"),
            ) or {}
        except Exception as exc:
            record["query_error"] = f"美客多商品详情读取失败：{str(exc)[:200]}"
    record["marketplace_item_id"] = marketplace_item_id
    record["_global_item_id"] = global_item_id or str(item_detail.get("cbt_item_id") or "").strip()
    record["_net_proceeds_usd"] = _extract_net_proceeds(item_detail)
    record["current_net_proceeds_usd"] = (
        str(record["_net_proceeds_usd"]) if record["_net_proceeds_usd"] is not None else ""
    )
    record["title"] = record.get("title") or str(item.get("title") or item_detail.get("title") or "")
    record["image_url"] = record.get("image_url") or str(
        item_detail.get("secure_thumbnail") or item_detail.get("thumbnail") or ""
    )
    record["category"] = record.get("category") or str(item.get("category_id") or "")
    record["query_status"] = "已读取"
    if record.get("query_error"):
        pass
    elif not marketplace_item_id and match_error:
        record["query_error"] = f"未能关联美客多商品：{match_error[:200]}"
    elif not record["_global_item_id"]:
        record["query_error"] = "商品没有可更新的 Global Selling 产品编号"
    elif not str(record.get("product_id") or "").strip():
        record["query_error"] = "缺少产品id，请先通过订单导入补全"
    return record


def _configured_workers(total: int, default: int, cap: int, env_name: str) -> int:
    try:
        requested = int(os.environ.get(env_name, default))
    except (TypeError, ValueError):
        requested = default
    return max(1, min(total, cap, requested))


def _run_read(task_id: str, file_bytes: bytes, filename: str, store_rows, allowed_token_ids):
    """Fill saved order metadata by order number without marketplace requests."""
    with _lock:
        task = _tasks.get(task_id)
        if not task:
            return
        task.update(status="preparing", message="正在按单号匹配订单文件", processed=0, total=0)
    try:
        uploaded = read_order_file(file_bytes, filename)
        by_number = {}
        for row in uploaded:
            number = str(row.get("order_number") or "").strip()
            if number not in by_number:
                by_number[number] = dict(row)
            else:
                # Repeated lines may contain complementary metadata.
                for key in HEADERS:
                    if row.get(key) and not by_number[number].get(key):
                        by_number[number][key] = row[key]
        duplicates = len(uploaded) - len(by_number)
        existing_ids = set(bit_db_api.get_weight_dimensions_record_order_numbers(list(by_number)) or [])
        supplements = [row for number, row in by_number.items() if number in existing_ids]
        unmatched = [row for number, row in by_number.items() if number not in existing_ids]
        enriched = 0
        for start in range(0, len(supplements), 100):
            saved = bit_db_api.save_weight_dimensions_records(
                supplements[start:start + 100], enrich_only=True,
            ) or {}
            enriched += int(saved.get("enriched") or 0)
        records = (bit_db_api.list_weight_dimensions_records(list(existing_ids)) or []) if existing_ids else []
        for row in unmatched:
            row["query_status"] = "未匹配"
            row["query_error"] = "已保存记录中没有此单号，未补充数据"
        records.extend(unmatched)
        records.sort(key=lambda row: str(row.get("time") or "").replace("T", " "), reverse=True)
    except Exception as exc:
        with _lock:
            task = _tasks.get(task_id)
            if task:
                task.update(
                    status="failed", finished_at=datetime.now().isoformat(timespec="seconds"),
                    message=f"订单文件补充失败：{str(exc)[:250]}",
                )
        return
    with _lock:
        task = _tasks.get(task_id)
        if task:
            task.update(
                records=records, status="ready", total=len(records), processed=len(records),
                finished_at=datetime.now().isoformat(timespec="seconds"),
                upload_enriched=enriched, upload_unmatched=len(unmatched), upload_skipped=duplicates,
                message=(f"补充完成，匹配已有订单 {len(supplements)} 条，补充 {enriched} 条，"
                         f"无需补充 {max(0, len(supplements) - enriched)} 条，"
                         f"未匹配单号 {len(unmatched)} 条，文件内重复单号 {duplicates} 条"),
            )


def start_upload(file_bytes: bytes, filename: str, owner: str,
                 store_rows: list[Mapping[str, Any]], allowed_token_ids: set[int] | None = None) -> dict[str, Any]:
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise ValueError("文件不能超过 16 MB")
    if not str(filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("请上传 .xlsx 或 .xlsm 订单文件")
    task_id = uuid.uuid4().hex
    with _lock:
        for old_id in list(_tasks):
            if _tasks[old_id].get("owner") == owner and _tasks[old_id].get("status") in {"ready", "failed"}:
                del _tasks[old_id]
        _tasks[task_id] = {
            "task_id": task_id, "owner": owner, "source": "upload", "status": "preparing",
            "message": "订单文件已接收，正在后台解析", "total": 0, "processed": 0,
            "records": [], "execute_status": "idle", "execute_message": "",
            "execute_processed": 0, "execute_total": 0, "created_at": datetime.now().isoformat(timespec="seconds"),
            "upload_skipped": 0,
            "authorized_ids": sorted(allowed_token_ids) if allowed_token_ids is not None else None,
            "agent_job_id": "",
        }
        _latest_task_by_owner[owner] = task_id
    thread = threading.Thread(
        target=_run_read, args=(task_id, file_bytes, filename, list(store_rows), allowed_token_ids),
        name=f"weight-dimensions-{task_id[:8]}", daemon=True,
    )
    thread.start()
    return {"task_id": task_id, "status": "preparing", "total": 0, "skipped": 0}


def _run_changed_read(task_id: str, owner: str, filters, store_rows, allowed_token_ids):
    """Load freight-change orders and fill package fields in the background."""
    try:
        _read_changed_records(task_id, owner, filters, store_rows, allowed_token_ids)
    except Exception as exc:
        with _lock:
            task = _tasks.get(task_id)
            if task:
                task.update(
                    status="failed", finished_at=datetime.now().isoformat(timespec="seconds"),
                    message=f"运费变更记录更新失败：{str(exc)[:250]}",
                )


def _read_changed_records(task_id, owner, filters, store_rows, allowed_token_ids):
    try:
        records = [dict(row or {}) for row in bit_db_api.list_weight_dimensions_changed_orders(filters) or []]
    except Exception as exc:
        with _lock:
            task = _tasks.get(task_id)
            if task:
                task.update(
                    status="failed", finished_at=datetime.now().isoformat(timespec="seconds"),
                    message=f"读取运费变更订单失败：{str(exc)[:250]}",
                )
        return

    # Keep page boundaries stable while background enrichment is running.
    records.sort(
        key=lambda row: (
            bool(str(row.get("time") or "").strip()),
            str(row.get("time") or "").strip().replace("T", " "),
        ),
        reverse=True,
    )
    with _lock:
        task = _tasks.get(task_id)
        if not task:
            return
        task.update(
            status="querying", message="正在从美客多官方补差 API 读取当前标记与实际重量尺寸",
            records=records, total=len(records), processed=0,
        )
    if not records:
        with _lock:
            task.update(
                status="ready", finished_at=datetime.now().isoformat(timespec="seconds"),
                message="没有找到官方确认由重量尺寸差异造成运费变动的订单",
            )
        return

    try:
        saved_records = []
        for start in range(0, len(records), 500):
            saved_records.extend(bit_db_api.list_weight_dimensions_records(
                [row.get("order_number") for row in records[start:start + 500]]
            ) or [])
    except Exception:
        saved_records = []
    saved_by_order = {
        str(row.get("order_number") or "").strip(): row
        for row in saved_records if isinstance(row, Mapping)
    }
    for record in records:
        saved = saved_by_order.get(str(record.get("order_number") or "").strip())
        if not saved:
            continue
        for key in ("execution_status", "execution_error", "execution_logs", "_zeshun_status", "_zying_status"):
            if key in saved:
                record[key] = saved[key]
        # Restore Excel metadata before querying and displaying refreshed orders.
        for key in ("time", "salesperson", "source", "company_store", "product_id",
                    "category", "title", "image_url", "tracking_number", "carrier", "region"):
            if not str(record.get(key) or "").strip() and saved.get(key):
                record[key] = saved[key]
        for key in ("actual_weight_g", "actual_dimensions_cm", "declared_weight_g", "declared_dimensions_cm"):
            if not record.get(key) and saved.get(key):
                record[key] = saved[key]

    store_map = _store_map(store_rows, allowed_token_ids)
    token_clients: dict[int, MercadoLibreClient] = {}
    token_errors: dict[int, str] = {}
    for record in records:
        token_id = int(record.get("_token_id") or record.get("store_token_id") or 0)
        record["_token_id"] = token_id or None
        try:
            if not token_id:
                raise ValueError("运费变更订单没有店铺授权")
            if token_id not in token_clients and token_id not in token_errors:
                token = bit_mysql.get_mercado_store_token(token_id)
                if not token:
                    raise ValueError("店铺授权不存在")
                token_clients[token_id], _ = _client_and_token(dict(token))
        except Exception as exc:
            token_errors[token_id] = str(exc)
            record["query_status"] = "查询失败"
            record["query_error"] = str(exc)[:300]

    worker_local = threading.local()
    progress_lock = threading.Lock()
    progress = sum(row.get("query_status") == "查询失败" for row in records)

    def read_and_count(record):
        nonlocal progress
        if record.get("query_status") != "查询失败":
            try:
                _read_one_record(record, store_map, token_clients, token_errors, worker_local)
            except Exception as exc:
                record["query_status"] = "查询失败"
                record["query_error"] = str(exc)[:300]
        with progress_lock:
            progress += 1
            with _lock:
                task = _tasks.get(task_id)
                if task:
                    task["processed"] = progress
                    task["message"] = f"已补全运费变更订单 {progress}/{len(records)} 条"

    # Commit each batch before starting the next so interrupted full syncs
    # retain their progress and database readers see results immediately.
    workers = _configured_workers(
        len(records), READ_MAX_WORKERS, READ_MAX_WORKERS,
        "WEIGHT_DIMENSIONS_READ_WORKERS",
    )
    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="wdr-change") as executor:
        for start in range(0, len(records), 100):
            batch = records[start:start + 100]
            futures = [executor.submit(read_and_count, row) for row in batch
                       if row.get("query_status") != "查询失败"]
            for future in as_completed(futures):
                future.result()
            # Persist failures too, for database-only diagnostics.
            bit_db_api.save_weight_dimensions_records(batch, refresh=True)

    records.sort(
        key=lambda row: (
            bool(str(row.get("time") or "").strip()),
            str(row.get("time") or "").strip().replace("T", " "),
        ),
        reverse=True,
    )
    failed_count = sum(row.get("query_status") == "查询失败" for row in records)
    with _lock:
        task = _tasks.get(task_id)
        if task:
            task.update(
                records=records, status="ready", processed=len(records),
                finished_at=datetime.now().isoformat(timespec="seconds"),
                message=f"运费变更订单读取完成，共 {len(records)} 条，失败 {failed_count} 条",
            )


def start_changed_refresh(owner: str, filters, store_rows, allowed_token_ids=None):
    task_id = uuid.uuid4().hex
    filters = dict(filters or {})
    if not filters.get("date_from") and not filters.get("date_to"):
        filters["date_from"] = (datetime.now() - timedelta(days=6)).strftime("%Y-%m-%d 00:00")
    authorized_ids = sorted(allowed_token_ids) if allowed_token_ids is not None else None
    with _lock:
        # Repeated refreshes should follow the running read, not start another
        # full database scan and thousands of duplicate marketplace requests.
        for task in _tasks.values():
            if (task.get("owner") == owner
                    and task.get("source") == "freight_changes"
                    and task.get("status") in {"preparing", "querying"}
                    and task.get("filters") == filters
                    and task.get("authorized_ids") == authorized_ids):
                return {"task_id": task["task_id"], "status": task["status"],
                        "total": task.get("total", 0)}
        for old_id in list(_tasks):
            if _tasks[old_id].get("owner") == owner and _tasks[old_id].get("source") == "freight_changes":
                if _tasks[old_id].get("status") in {"ready", "failed"}:
                    del _tasks[old_id]
        _tasks[task_id] = {
            "task_id": task_id, "owner": owner, "source": "freight_changes",
            "status": "preparing", "message": "正在读取运费变更订单",
            "total": 0, "processed": 0, "records": [], "execute_status": "idle",
            "execute_message": "", "execute_processed": 0, "execute_total": 0,
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "filters": filters, "authorized_ids": authorized_ids, "agent_job_id": "",
        }
        _latest_task_by_owner[owner] = task_id
    thread = threading.Thread(
        target=_run_changed_read,
        args=(task_id, owner, dict(filters or {}), list(store_rows), allowed_token_ids),
        name=f"weight-dimensions-changed-{task_id[:8]}", daemon=True,
    )
    thread.start()
    return {"task_id": task_id, "status": "preparing", "total": 0}


def load_saved_changes(owner: str, filters, allowed_token_ids=None, *, page_size=50):
    """Build an executable view from persisted records; never query marketplaces."""
    filters = dict(filters or {})
    if not filters.get("date_from") and not filters.get("date_to"):
        filters["date_from"] = (datetime.now() - timedelta(days=6)).strftime("%Y-%m-%d 00:00")
    if allowed_token_ids is not None:
        requested = filters.get("store_ids", allowed_token_ids)
        filters["store_ids"] = sorted(set(requested) & set(allowed_token_ids))
    first_page = bit_db_api.list_weight_dimensions_records(filters=filters, page=1, page_size=page_size)
    total = first_page["record_total"]
    task_id = uuid.uuid4().hex
    now = datetime.now().isoformat(timespec="seconds")
    with _lock:
        for old_id, task in list(_tasks.items()):
            if (task.get("owner") == owner and task.get("source") == "saved_changes"
                    and task.get("execute_status", "idle") not in {"running", "queued", "stopping"}):
                del _tasks[old_id]
        _tasks[task_id] = {
            "task_id": task_id, "owner": owner, "source": "saved_changes",
            "status": "ready", "records": [], "total": total, "processed": total,
            "database_view": True,
            "message": f"已从数据库读取 {total} 条记录；后台每天 05:00 更新",
            "execute_status": "idle", "execute_message": "", "execute_processed": 0,
            "execute_total": 0, "created_at": now, "finished_at": now,
            "filters": filters, "agent_job_id": "",
            "authorized_ids": sorted(allowed_token_ids) if allowed_token_ids is not None else None,
        }
        _latest_task_by_owner[owner] = task_id
    return {"task_id": task_id, "status": "ready", "total": total,
            **first_page, "records": [_public_record(row) for row in first_page["records"]]}


def start_full_refresh(owner: str, store_rows, allowed_token_ids=None):
    """Refresh every freight-change record, without relying on an upload file."""
    return start_changed_refresh(
        owner,
        {"date_from": "2000-01-01 00:00", "date_to": "2099-12-31 23:59",
         "store_ids": sorted(allowed_token_ids or [])},
        store_rows,
        allowed_token_ids,
    )


def _get_owned_task(task_id: str, owner: str):
    with _lock:
        task = _tasks.get(str(task_id or ""))
        if not task or task.get("owner") != owner:
            raise KeyError("查询任务不存在或已过期")
        return task


def _public_record(row: Mapping[str, Any]) -> dict[str, Any]:
    result = {key: row.get(key, "") for key, _label in PUBLIC_COLUMNS}
    result["measurement_notes"] = row.get("measurement_notes", {})
    result["execution_status"] = row.get("execution_status", "")
    result["zeshun_execution_status"] = row.get("_zeshun_status") or "未执行"
    result["zying_execution_status"] = row.get("_zying_status") or "未执行"
    result["execution_error"] = row.get("execution_error", "")
    result["execution_logs"] = row.get("execution_logs", [])
    net_value = row.get("current_net_proceeds_usd")
    if net_value in (None, ""):
        net_value = row.get("_net_proceeds_usd")
    result["current_net_proceeds_usd"] = str(net_value) if net_value is not None else ""
    can_zying = bool(
        row.get("_dimension_compensation_confirmed")
        and
        str(row.get("product_id") or "").strip()
        and str(row.get("actual_weight_g") or "").strip()
        and row.get("_zying_status") != "成功"
    )
    can_zeshun = bool(
        row.get("_dimension_compensation_confirmed")
        and row.get("query_status") == "已读取"
        and str(row.get("actual_weight_g") or "").strip()
        and row.get("_token_id")
        and row.get("_global_item_id")
        and row.get("marketplace_item_id")
        and row.get("_zeshun_status") != "成功"
    )
    result["can_execute_zying"] = can_zying
    result["can_execute_zeshun"] = can_zeshun
    result["can_execute"] = can_zying or can_zeshun
    return result


def _zeshun_block_reason(row: Mapping[str, Any]) -> str:
    """Explain why a freight record cannot enter the Zeshun update pipeline."""
    if row.get("_zeshun_status") == "成功":
        return ""
    if not row.get("_dimension_compensation_confirmed"):
        return "运费变动记录尚未完成重量尺寸核对"
    if row.get("query_status") != "已读取":
        return str(row.get("query_error") or "订单商品信息尚未读取成功")
    if not str(row.get("actual_weight_g") or "").strip():
        return "缺少经官方核验的实际重量"
    if not row.get("_token_id"):
        return "缺少订单关联店铺"
    if not str(row.get("_global_item_id") or "").strip():
        return str(row.get("query_error") or "商品没有可更新的 Global Selling 产品编号")
    if not str(row.get("marketplace_item_id") or "").strip():
        return str(row.get("query_error") or "商品没有可更新的站点链接编号")
    return "当前记录不满足泽顺数据更新条件"


def _finish_stopped_execute(task):
    action = task.get("execute_action")
    selected = set(task.get("execute_order_numbers") or ())
    for row in task.get("records") or ():
        if selected and str(row.get("order_number") or "") not in selected:
            continue
        pending = False
        for key in (["_zying_status"] if action == "zying" else
                    ["_zeshun_status"] if action == "zeshun" else ["_zying_status", "_zeshun_status"]):
            if row.get(key) not in {"成功", "失败", "跳过", "已终结"}:
                row[key] = "已终结"
                pending = True
        if pending:
            row["execution_status"] = "已终结"
            row["execution_error"] = "任务已终结，未执行的更新可重新选择后执行"
            _log_execution(row, "任务", "已终结", row["execution_error"])
    with _lock:
        task["execute_status"] = "stopped"
        task["execute_message"] = "任务已终结，已完成结果已保留，未执行记录可重新选择后执行"


def stop_execute(task_id: str, owner: str, *, agent_cancel=None):
    task = _get_owned_task(task_id, owner)
    with _lock:
        if task.get("execute_status") not in {"queued", "running", "stopping"}:
            return {"status": task.get("execute_status", "idle"), "message": task.get("execute_message", "当前没有执行中的任务")}
        previous_status = task["execute_status"]
        job_id = task.get("agent_job_id")
        if job_id and agent_cancel is None:
            raise ValueError("未配置本机 Agent 任务终结接口")
        task["execute_stop_requested"] = True
        task["execute_status"] = "stopping"
        task["execute_message"] = "正在终结任务，当前产品处理结束后停止后续更新"
        task["agent_cancel"] = agent_cancel
    if job_id:
        try:
            agent_cancel(job_id)
        except Exception:
            with _lock:
                if task.get("execute_status") == "stopping":
                    task["execute_status"] = previous_status
                    task["execute_stop_requested"] = False
            raise
    return {"status": "stopping", "message": task["execute_message"]}


def shutdown_executions(timeout=30):
    """Stop server executions and allow current writes to finish before exit."""
    global _executions_shutting_down
    deadline = time.monotonic() + max(0, float(timeout))
    with _lock:
        _executions_shutting_down = True
        workers = list(_execution_threads.items())
        for task_id, _thread in workers:
            task = _tasks.get(task_id)
            if task and task.get("execute_status") in {"queued", "running", "stopping"}:
                task["execute_stop_requested"] = True
                task["execute_status"] = "stopping"
                task["execute_message"] = "服务正在关闭，当前更新收尾后终结任务"
    current = threading.current_thread()
    for _task_id, worker in workers:
        if worker is not current:
            worker.join(max(0, deadline - time.monotonic()))
    pending = sum(worker.is_alive() for _task_id, worker in workers)
    if pending:
        _logger.warning("服务关闭时仍有 %s 个重量尺寸执行线程未收尾", pending)
    return {"requested": len(workers), "pending": pending}


def _apply_agent_job(task, job):
    """Mirror Agent queue, execution, and result states into the WDR task."""
    if not task or not job:
        return
    job_status = str(job.get("status") or "").strip().lower()
    result = job.get("result") or {}
    result_rows = result.get("rows") if isinstance(result, Mapping) else []
    result_by_order = {
        str(row.get("order_number") or "").strip(): row
        for row in result_rows or () if isinstance(row, Mapping)
    }
    records = task.get("records") or []
    selected_orders = {
        str(value or "").strip()
        for value in task.get("execute_order_numbers") or ()
        if str(value or "").strip()
    }
    if job_status in {"queued", "running", "stopping"}:
        task["execute_status"] = "stopping" if job_status == "stopping" or task.get("execute_stop_requested") else "queued" if job_status == "queued" else "running"
        task["execute_message"] = job.get("message") or "Agent 正在逐个更新智赢产品"
        row_status = "排队中" if job_status == "queued" else "执行中"
        progress_message = (
            "任务已提交本机 Agent 队列，尚未开始更新"
            if job_status == "queued" else "本机 Agent 已领取任务，正在更新智赢产品"
        )
        for row in records:
            order_number = str(row.get("order_number") or "").strip()
            if selected_orders and order_number not in selected_orders:
                continue
            if row.get("execution_status") == row_status and row.get("_zying_status") == row_status:
                continue
            row["_zying_status"] = row_status
            row["execution_status"] = row_status
            row["execution_error"] = ""
            _log_execution(row, "智赢 Agent", row_status, progress_message)
        return
    if task.get("agent_applied_status") == job_status and task.get("agent_applied_result"):
        return
    for row in records:
        order_number = str(row.get("order_number") or "").strip()
        if selected_orders and order_number not in selected_orders:
            continue
        item = result_by_order.get(order_number) or {}
        success = str(item.get("status") or "").lower() in {"success", "成功"}
        if success:
            row["_zying_status"] = "成功"
            row["execution_status"] = "智赢产品已更新"
            row["execution_error"] = ""
            _log_execution(
                row, "智赢 Agent", "成功",
                str(item.get("message") or "重量、尺寸和产品级别“重点”已保存并回读确认"),
            )
        else:
            message = str(
                item.get("message")
                or job.get("message")
                or "Agent 未返回该订单的更新结果"
            )
            stopped = job_status == "stopped" and not item
            row["_zying_status"] = "已终结" if stopped else "失败"
            row["execution_status"] = row["_zying_status"]
            row["execution_error"] = "任务已终结，未返回该产品的保存确认，可重新执行核对" if stopped else f"智赢产品更新失败：{message[:250]}"
            _log_execution(row, "智赢 Agent", row["_zying_status"], row["execution_error"])
    task["agent_applied_status"] = job_status
    task["agent_applied_result"] = True
    task["execute_status"] = "stopped" if job_status == "stopped" else "completed"
    task["execute_processed"] = len(result_by_order) if job_status == "stopped" else task.get("execute_total") or len(records)
    succeeded = sum(
        row.get("_zying_status") == "成功" and str(row.get("order_number") or "") in selected_orders
        for row in records
    )
    task["execute_message"] = f"智赢 Agent {'任务已终结' if job_status == 'stopped' else '执行结束'}，成功 {succeeded} 条"


def task_status(task_id: str, owner: str, agent_job_provider=None, *, page=None, page_size=50, include_records=True) -> dict[str, Any]:
    task = _get_owned_task(task_id, owner)
    agent_job_id = str(task.get("agent_job_id") or "").strip()
    if agent_job_id and agent_job_provider:
        try:
            job = agent_job_provider(agent_job_id)
            # Applying Agent state transitions writes execution logs to the
            # database. Do not hold the task lock across that I/O.
            _apply_agent_job(task, job)
        except Exception as exc:
            with _lock:
                task["execute_message"] = f"读取 Agent 状态失败：{str(exc)[:250]}"
    database_page = None
    if task.get("database_view") and include_records:
        database_page = bit_db_api.list_weight_dimensions_records(
            filters=task["filters"], page=page or 1, page_size=page_size,
        )
    with _lock:
        rows = task["records"]
        if task.get("source") in {"upload", "freight_changes"}:
            rows = [row for row in rows if row.get("_dimension_compensation_confirmed")]
        pagination = {}
        if database_page is not None:
            rows = database_page["records"]
            pagination = {key: database_page[key] for key in ("page", "page_size", "record_total")}
            task["total"] = database_page["record_total"]
            active_rows = {row["order_number"]: row for row in task["records"]}
            rows = [active_rows.get(row["order_number"], row) for row in rows]
        elif page is not None:
            page_size = max(1, min(100, int(page_size)))
            page_count = max(1, (len(rows) + page_size - 1) // page_size)
            page = max(1, min(int(page), page_count))
            pagination = {"page": page, "page_size": page_size, "record_total": len(rows)}
            rows = rows[(page - 1) * page_size:page * page_size]
        return {
            **pagination,
            "task_id": task["task_id"], "status": task["status"],
            "message": task.get("message", ""), "total": task.get("total", 0),
            "processed": task.get("processed", 0), "records": [_public_record(row) for row in rows] if include_records else [],
            "created_at": task.get("created_at", ""), "finished_at": task.get("finished_at", ""),
            "execute_status": task.get("execute_status", "idle"),
            "execute_message": task.get("execute_message", ""),
            "execute_processed": task.get("execute_processed", 0),
            "execute_total": task.get("execute_total", 0),
            "upload_skipped": task.get("upload_skipped", 0),
            "source": task.get("source", "upload"),
            "agent_job_id": agent_job_id,
        }


def latest_task_status(owner: str, agent_job_provider=None) -> dict[str, Any] | None:
    with _lock:
        task_id = _latest_task_by_owner.get(str(owner or ""))
    if not task_id:
        return None
    try:
        return task_status(task_id, owner, agent_job_provider=agent_job_provider)
    except KeyError:
        return None


def start_latest_execute(owner: str, erp_browser_factory=None, erp_browser_preflight=None) -> dict[str, Any]:
    with _lock:
        task_id = _latest_task_by_owner.get(str(owner or ""))
    if not task_id:
        raise ValueError("请先在控制台上传订单文件并完成查询")
    # The browser-extension endpoint retains its legacy combined workflow;
    # the console page uses the two explicit actions instead.
    return start_execute(
        task_id, owner, action="combined", erp_browser_factory=erp_browser_factory,
        erp_browser_preflight=erp_browser_preflight,
    )


def export_xlsx(task_id: str, owner: str) -> bytes:
    task = _get_owned_task(task_id, owner)
    if task.get("status") != "ready":
        raise ValueError("查询完成后才能导出")
    rows = bit_db_api.list_weight_dimensions_records(filters=task["filters"]) if task.get("database_view") else task["records"]
    if task.get("source") in {"upload", "freight_changes"}:
        rows = [row for row in rows if row.get("_dimension_compensation_confirmed")]
    return export_records_xlsx(rows)


def export_records_xlsx(records) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "重量尺寸变更记录"
    labels = [label for _key, label in EXPORT_COLUMNS]
    sheet.append(labels)
    for row in records or []:
        values = []
        for key, _label in EXPORT_COLUMNS:
            if key == "current_net_proceeds_usd":
                value = row.get("current_net_proceeds_usd", row.get("_net_proceeds_usd"))
                values.append(str(value) if value is not None else "")
            elif key == "execution_logs":
                values.append("\n".join(
                    f"{log.get('time', '')} [{log.get('stage', '')}/{log.get('status', '')}] "
                    + " · ".join(
                        str(value) for value in (
                            log.get("site_name") or log.get("site_id"),
                            log.get("link_name"), log.get("marketplace_item_id"),
                            log.get("salesperson"), log.get("store_group"), log.get("store_name"),
                            log.get("message"),
                        ) if value not in (None, "")
                    )
                    for log in row.get("execution_logs", []) if isinstance(log, Mapping)
                ))
            else:
                values.append(row.get(key, ""))
        sheet.append(values)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [20, 20, 24, 20, 16, 24, 28, 16, 18, 46, 20, 26, 18, 16, 18, 22, 18, 18, 22, 18, 22, 16, 44, 22, 18, 44, 44, 64, 64]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="203E5B")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@" if cell.column in {1, 3, 7, 10, 11, 20} else "General"
    for index, (key, _) in enumerate(EXPORT_COLUMNS, start=1):
        color = "FFF3DF" if key.startswith("declared_") else "E5F4EF" if key.startswith("actual_") else None
        if color:
            for cells in sheet.iter_rows(min_col=index, max_col=index):
                cells[0].fill = PatternFill("solid", fgColor=color)
                if cells[0].row == 1:
                    cells[0].font = Font(bold=True, color="263A50")
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _package_value(attribute, unit):
    """Normalize API package attributes for read-back comparison."""
    raw = attribute.get("value_struct") or {}
    if not raw:
        match = re.fullmatch(r"\s*([\d.]+)\s*(\w+)\s*", str(attribute.get("value_name") or ""))
        if not match:
            return None
        raw = {"number": match[1], "unit": match[2]}
    factors = {"g": {"g": "1", "kg": "1000", "lb": "453.59237", "oz": "28.349523125"},
               "cm": {"cm": "1", "mm": "0.1", "m": "100", "in": "2.54"}}
    try:
        result = Decimal(str(raw["number"])) * Decimal(factors[unit][raw["unit"]])
        return result if result.is_finite() and result > 0 else None
    except (KeyError, InvalidOperation, TypeError):
        return None


def _update_error(exc):
    """Keep the useful API cause instead of truncating before its description."""
    message = str(exc)
    try:
        body = json.loads(message[message.index("{"):])
        causes = body.get("cause") or []
        detail = "; ".join(str(c.get("code", "")) + ": " +
                           str(c.get("message") or c.get("description") or "")
                           for c in causes if isinstance(c, Mapping))
        if detail:
            return (str(body.get("message") or "更新失败") + "；" + detail)[:1600]
    except (ValueError, TypeError, AttributeError):
        pass
    return message[:1600]


def _marketplace_package_weight_g(weight_g) -> str:
    """Use Mercado Libre's integer grams without rounding invalid input up."""
    try:
        weight = Decimal(str(weight_g))
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise ValueError("美客多要求包装重量至少 50 克") from exc
    if not weight.is_finite() or weight < 50:
        raise ValueError("美客多要求包装重量至少 50 克")
    return format(weight.to_integral_value(rounding=ROUND_HALF_UP), "f")


def _package_attributes(client: MercadoLibreClient, global_item_id: str,
                        weight_g: str, dimensions_cm: str = "") -> None:
    weight = Decimal(_marketplace_package_weight_g(weight_g))
    values = {"PACKAGE_WEIGHT": (weight, "g")}
    pieces = []
    if str(dimensions_cm or "").strip():
        pieces = [Decimal(part) for part in str(dimensions_cm).lower().replace("×", "x").split("x")]
        if len(pieces) != 3:
            raise ValueError("实际尺寸格式应为 长x宽x高")
        values.update({
            "PACKAGE_LENGTH": (pieces[0], "cm"),
            "PACKAGE_WIDTH": (pieces[1], "cm"),
            "PACKAGE_HEIGHT": (pieces[2], "cm"),
        })
    for number, unit in values.values():
        if not number.is_finite() or number <= 0:
            raise ValueError("重量或尺寸必须是大于 0 的有效数值")
    if any(value < 3 for value in pieces):
        raise ValueError("美客多要求包装长宽高至少 3 厘米")
    # Global Selling reads use /marketplace/items; /global/items is for writes.
    remote = client.get_marketplace_item(global_item_id)
    attrs = {a["id"]: a for a in remote.get("attributes") or [] if isinstance(a, Mapping) and a.get("id")}
    # The API needs all four package fields, even for a weight-only update.
    # Preserve existing dimensions, but never copy stale values/value_struct or
    # unrelated immutable attributes into the write payload.
    for key, (_, unit) in _PACKAGE_ATTRIBUTES.items():
        if key not in values:
            current = _package_value(attrs.get(key, {}), unit)
            if current is None:
                raise ValueError(f"官方未提供尺寸，且商品缺少有效 {key}，无法保留原尺寸")
            values[key] = (current, unit)

    def differences(item):
        actual = {a["id"]: a for a in item.get("attributes") or [] if isinstance(a, Mapping) and a.get("id")}
        return [key for key, (number, unit) in values.items()
                if _package_value(actual.get(key, {}), unit) != number]

    if not differences(remote):
        return
    updates = [{"id": key, "value_name": f"{format(number.normalize(), 'f')} {unit}"}
               for key, (number, unit) in values.items()]
    # A user_product_id alone does not imply the account exposes the newer UP
    # endpoint. The documented global item update remains compatible; sending
    # only clean package fields avoids legacy metadata conflicts on both models.
    client.update_global_item(global_item_id, {"attributes": updates})
    for delay in (0, 1, 2, 4, 8, 8):
        if delay:
            time.sleep(delay)
        missing = differences(client.get_marketplace_item(global_item_id))
        if not missing:
            return
    raise RuntimeError("提交后回读尚未确认重量尺寸生效：" + ", ".join(missing))


def _closed_listing_reason(item):
    if "deleted" in (item.get("sub_status") or []):
        return "美客多链接已删除，跳过重量尺寸及净收益更新"
    if str(item.get("status") or "").lower() == "closed":
        return "美客多链接已关闭，跳过重量尺寸及净收益更新"
    return ""


def _normalize_history_site_id(value):
    text = str(value or "").strip().upper()
    if text in _SITE_NAME_BY_ID:
        return text
    country_sites = {country.upper(): site_id for country, site_id in _COUNTRY_TO_SITE.items()}
    country_sites.update({
        "MX": "MLM", "MEXICO": "MLM", "BR": "MLB", "BRAZIL": "MLB",
        "CL": "MLC", "CHILE": "MLC", "CO": "MCO", "COLOMBIA": "MCO",
        "AR": "MLA", "ARGENTINA": "MLA", "UY": "MLU", "URUGUAY": "MLU",
        "PE": "MPE", "PERU": "MPE", "EC": "MEC", "ECUADOR": "MEC",
    })
    return country_sites.get(text, "")


def _history_site_settings(token_id):
    payload = bit_db_api.list_mercado_store_site_settings(token_id)
    rows = payload.get("rows", []) if isinstance(payload, Mapping) else payload
    return [row for row in (rows or []) if isinstance(row, Mapping)]


def _order_history_context(row, settings_cache=None):
    """Capture the order's store/site ownership for product-level update logs."""
    try:
        token_id = int(row.get("_token_id") or row.get("store_token_id") or 0)
    except (TypeError, ValueError):
        token_id = 0
    site_id = (
        _normalize_history_site_id(row.get("site_id"))
        or _normalize_history_site_id(row.get("region"))
        or _normalize_history_site_id(str(row.get("marketplace_item_id") or "")[:3])
    )
    salesperson = str(row.get("salesperson") or "").strip()
    store_group = str(row.get("store_group") or row.get("group_name") or "").strip()
    if token_id > 0 and site_id and settings_cache is not None:
        if token_id not in settings_cache:
            try:
                settings_cache[token_id] = _history_site_settings(token_id)
            except Exception:
                settings_cache[token_id] = []
        setting = next((value for value in settings_cache[token_id]
                        if _normalize_history_site_id(value.get("site_id")) == site_id), {})
        salesperson = salesperson or str(setting.get("salesperson") or "").strip()
        store_group = store_group or str(setting.get("group_name") or "").strip()
    return {
        "site_id": site_id,
        "site_name": _SITE_NAME_BY_ID.get(site_id, ""),
        "link_name": str(row.get("title") or "").strip(),
        "salesperson": salesperson,
        "store_group": store_group,
        "store_name": str(row.get("company_store") or row.get("store_name") or "").strip(),
        "marketplace_item_id": str(row.get("marketplace_item_id") or "").strip(),
        "product_id": str(row.get("product_id") or "").strip(),
    }


def _listing_history_context(row, token_id, item_id, link, remote, settings_cache):
    link = link if isinstance(link, Mapping) else {}
    remote = remote if isinstance(remote, Mapping) else {}
    site_id = _normalize_history_site_id(
        remote.get("site_id") or link.get("site_id") or str(item_id or "")[:3]
    )
    context = _order_history_context(row, settings_cache)
    context.update({
        "site_id": site_id,
        "site_name": _SITE_NAME_BY_ID.get(site_id, ""),
        "link_name": str(remote.get("title") or link.get("title") or row.get("title") or "").strip(),
        "store_name": str(link.get("store_name") or row.get("company_store") or "").strip(),
        "marketplace_item_id": str(item_id or "").strip(),
        "product_id": str(row.get("product_id") or "").strip(),
    })
    if token_id and site_id:
        token_id = int(token_id)
        if token_id not in settings_cache:
            try:
                settings_cache[token_id] = _history_site_settings(token_id)
            except Exception:
                settings_cache[token_id] = []
        setting = next((value for value in settings_cache[token_id]
                        if _normalize_history_site_id(value.get("site_id")) == site_id), {})
        context["salesperson"] = str(setting.get("salesperson") or link.get("salesperson") or context.get("salesperson") or "").strip()
        context["store_group"] = str(setting.get("group_name") or link.get("group_name") or context.get("store_group") or "").strip()
    return context


def _expand_same_title_records(records, authorized_ids, settings_cache=None, *, stop_requested=None):
    """Resolve every current exact-title listing before writing any attributes."""
    clients, details, matches, lookup_errors, token_errors = {}, {}, {}, {}, {}
    link_rows, settings_cache = {}, settings_cache if settings_cache is not None else {}
    expanded = []

    def check_stopped():
        if stop_requested is not None and stop_requested():
            raise _ExecutionStopped()

    def detail(token_id, item_id):
        check_stopped()
        key = (token_id, item_id)
        if token_id in token_errors:
            raise ValueError(token_errors[token_id])
        if key in lookup_errors:
            raise ValueError(lookup_errors[key])
        if key not in details:
            if token_id not in clients:
                try:
                    token = bit_mysql.get_mercado_store_token(token_id)
                    clients[token_id], _ = _client_and_token(dict(token or {}))
                except Exception as exc:
                    token_errors[token_id] = _update_error(exc)
                    raise
            try:
                details[key] = clients[token_id].get_marketplace_item(
                    item_id, attributes=("id", "title", "cbt_item_id", "status", "sub_status")
                )
            except Exception as exc:
                lookup_errors[key] = _update_error(exc)
                raise
        return details[key]

    for row in records:
        if stop_requested is not None and stop_requested():
            break
        if not _public_record(row)["can_execute_zeshun"]:
            if row.get("_zeshun_status") == "成功":
                continue
            row["_zeshun_status"] = "失败"
            row["execution_status"] = "失败"
            row["execution_error"] = _zeshun_block_reason(row)
            _log_execution(row, "准备", "失败", row["execution_error"])
            continue
        row["execution_error"] = ""
        row["execution_status"] = "执行中"
        row["_zeshun_status"] = "执行中"
        try:
            token_id = int(row["_token_id"])
            item_id = str(row["marketplace_item_id"])
            if authorized_ids is not None and token_id not in authorized_ids:
                raise ValueError("订单关联店铺不在授权范围内")
            original = detail(token_id, item_id)
            check_stopped()
            title = str(original.get("title") or "").strip()
            if not title:
                raise ValueError("无法读取关联链接标题，不能确定同名链接范围")
            row["_history_context"] = _listing_history_context(
                row, token_id, item_id,
                {"token_id": token_id, "item_id": item_id, "site_id": original.get("site_id") or item_id[:3],
                 "title": title, "store_name": row.get("company_store") or ""},
                original, settings_cache,
            )
            if title not in matches:
                found = {}
                page = 1
                while True:
                    check_stopped()
                    result = bit_db_api.list_mercado_store_links(
                        search=title, token_ids=authorized_ids, include_categories=False,
                        sort_by="item_id", sort_order="asc", page=page, page_size=1000,
                    )
                    for link in result["rows"]:
                        check_stopped()
                        link_token = int(link["token_id"])
                        if authorized_ids is not None and link_token not in authorized_ids:
                            continue
                        if str(link.get("title") or "").strip() != title:
                            continue
                        key = (link_token, str(link["item_id"]))
                        link_rows[key] = link
                        try:
                            remote = detail(*key)
                            check_stopped()
                            if str(remote.get("title") or "").strip() != title:
                                continue
                            skip_reason = _closed_listing_reason(remote)
                            if skip_reason:
                                found[key] = ("", "", skip_reason)
                                continue
                            global_id = str(remote.get("cbt_item_id") or "").strip()
                            if not global_id:
                                raise ValueError("缺少 Global 商品关联")
                            found[key] = (global_id, "", "")
                        except _ExecutionStopped:
                            raise
                        except Exception as exc:
                            # An inaccessible sibling must not discard every
                            # valid listing in this title group. Retain failure
                            # children so the final summary cannot claim success.
                            found[key] = ("", _update_error(exc), "")
                    if page >= int(result["pages"] or 1):
                        break
                    page += 1
                matches[title] = found
            targets = dict(matches[title])
            original_global = str(original.get("cbt_item_id") or row.get("_global_item_id") or "")
            original_skip = _closed_listing_reason(original)
            targets[(token_id, item_id)] = (original_global, "" if original_global or original_skip else "缺少 Global 商品关联", original_skip)
            link_rows.setdefault((token_id, item_id), {
                "token_id": token_id, "item_id": item_id,
                "site_id": original.get("site_id") or item_id[:3],
                "title": title, "store_name": row.get("company_store") or "",
            })
            for (target_token, target_item), (global_id, lookup_error, skip_reason) in targets.items():
                check_stopped()
                child = dict(row)
                for field in ("_package_status", "_net_status", "_zeshun_status", "_expansion_error", "_expansion_skip"):
                    child.pop(field, None)
                child.update(_token_id=target_token, marketplace_item_id=target_item,
                             _global_item_id=global_id, _execution_parent=row,
                             execution_logs=[])
                child["_history_context"] = _listing_history_context(
                    row, target_token, target_item, link_rows.get((target_token, target_item)),
                    details.get((target_token, target_item)), settings_cache,
                )
                if lookup_error:
                    child.update(_zeshun_status="失败", _expansion_error=lookup_error,
                                 execution_status="失败", execution_error=lookup_error)
                    _log_execution(child, "同名链接", "失败", lookup_error)
                elif skip_reason:
                    child.update(_zeshun_status="跳过", _expansion_skip=skip_reason,
                                 execution_status="跳过", execution_error="")
                    _log_execution(child, "同名链接", "跳过", skip_reason)
                expanded.append(child)
            skipped_count = sum(bool(value[2]) for value in targets.values())
            _log_execution(row, "同名链接", "进行中", f"标题“{title}”匹配 {len(targets)} 条链接（含订单关联链接），"
                           + (f"其中 {skipped_count} 条已关闭或删除并跳过；" if skipped_count else "")
                           + "其余可更新链接将逐条保留并重提各自净收益")
        except _ExecutionStopped:
            break
        except Exception as exc:
            row["_zeshun_status"] = "失败"
            row["execution_status"] = "失败"
            row["execution_error"] = f"同名链接查询失败：{_update_error(exc)}"
            _log_execution(row, "同名链接", "失败", row["execution_error"])
    return expanded


def _finish_same_title_records(originals, expanded, *, stopped=False):
    by_parent = defaultdict(list)
    for row in expanded:
        by_parent[id(row["_execution_parent"])].append(row)
    for original in originals:
        children = by_parent[id(original)]
        if not children:
            continue
        skipped = sum(bool(row.get("_expansion_skip")) for row in children)
        targets = [row for row in children if not row.get("_expansion_skip")]
        if stopped:
            unfinished = [row for row in targets
                          if row.get("_zeshun_status") not in {"成功", "失败", "跳过", "已终结"}]
            for child in unfinished:
                child.update(_zeshun_status="已终结", execution_status="已终结",
                             execution_error="任务已终结，未完成的链接更新可重新执行核对")
                _log_execution(child, "任务", "已终结", child["execution_error"])
        succeeded = sum(row.get("_zeshun_status") == "成功" for row in targets)
        ended = sum(row.get("_zeshun_status") == "已终结" for row in targets)
        complete = bool(targets) and succeeded == len(targets)
        failures = [f"{row['marketplace_item_id']}：{row.get('execution_error') or '未完成'}"
                    for row in targets if row.get("_zeshun_status") not in {"成功", "已终结"}]
        status = "成功" if complete else "失败" if failures else "已终结" if ended else "跳过"
        original["_zeshun_status"] = status
        original["execution_status"] = "完成" if complete else "部分完成" if failures else status
        original["execution_error"] = "；".join(failures) or (
            "任务已终结，未完成的链接更新可重新执行核对" if ended else "")
        for child in children:
            if (child["_token_id"], child["marketplace_item_id"]) == (original["_token_id"], original["marketplace_item_id"]):
                original["current_net_proceeds_usd"] = child.get("current_net_proceeds_usd", "")
        _log_execution(original, "同名链接", status,
                       f"同名链接更新成功 {succeeded}/{len(targets)} 条"
                       + (f"；跳过已关闭或删除链接 {skipped} 条" if skipped else "")
                       + (f"；任务终结，未完成 {ended} 条" if ended else "")
                       + (f"；{original['execution_error']}" if failures else ""))


def _log_execution(row: dict[str, Any], stage: str, status: str, message: str) -> None:
    context = row
    parent = row.get("_execution_parent")
    if parent is not None:
        message = f"链接 {row['marketplace_item_id']}：{message}"
        row = parent
    entry = {
        "time": datetime.now().isoformat(timespec="seconds"),
        "stage": str(stage), "status": str(status), "message": str(message)[:2000],
        "order_number": str(row.get("order_number") or ""),
        "execution_id": str(context.get("_execution_id") or ""),
        "action": str(context.get("_execution_action") or ""),
        "operator": str(context.get("_execution_operator") or ""),
    }
    history_context = context.get("_history_context") or row.get("_history_context") or {}
    entry.update({
        "site_id": str(history_context.get("site_id") or ""),
        "site_name": str(history_context.get("site_name") or ""),
        "link_name": str(history_context.get("link_name") or ""),
        "salesperson": str(history_context.get("salesperson") or ""),
        "store_group": str(history_context.get("store_group") or ""),
        "store_name": str(history_context.get("store_name") or ""),
        "product_id": str(history_context.get("product_id") or ""),
    })
    if history_context.get("marketplace_item_id"):
        entry["marketplace_item_id"] = str(history_context["marketplace_item_id"])
    # Snapshot the actual child target and submitted values. Never infer history
    # from mutable order measurements when the page is read later.
    if parent is not None or stage in {"美客多链接", "净收益"}:
        entry.update(
            marketplace_item_id=str(context.get("marketplace_item_id") or ""),
            global_item_id=str(context.get("_global_item_id") or ""),
            token_id=context.get("_token_id"),
            submitted_weight_g=str(context.get("_submitted_weight_g") or ""),
            submitted_dimensions_cm=str(context.get("_submitted_dimensions_cm") or ""),
            dimensions_preserved=bool(context.get("_dimensions_preserved")),
        )
    if stage == "净收益" and context.get("current_net_proceeds_usd") is not None:
        entry["net_proceeds_usd"] = str(context["current_net_proceeds_usd"])
    if context.get("_automation_run_id"):
        from bit.automation_center import journal
        journal(context["_automation_run_id"], entry)
    row.setdefault("execution_logs", []).append(entry)
    updates = {
        "execution_status": row.get("execution_status", ""),
        "execution_error": row.get("execution_error", ""),
    }
    if row.get("current_net_proceeds_usd") is not None:
        updates["current_net_proceeds_usd"] = str(row["current_net_proceeds_usd"])
    for key in ("_zeshun_status", "_zying_status"):
        if key in row:
            updates[key] = row[key]
    try:
        bit_db_api.append_weight_dimensions_record_log(row.get("order_number"), entry, updates)
    except Exception as exc:
        row["execution_logs"].append({
            "time": datetime.now().isoformat(timespec="seconds"),
            "stage": "日志", "status": "保存失败", "message": str(exc)[:300],
        })


def _run_execute(
    task_id: str,
    owner: str,
    action="combined",
    selected_order_numbers=None,
    erp_browser_factory=None,
):
    task = _get_owned_task(task_id, owner)
    with _lock:
        all_records = task["records"]
        selected = {str(value or "").strip() for value in selected_order_numbers or () if str(value or "").strip()}
        records = [
            row for row in all_records
            if not selected or str(row.get("order_number") or "").strip() in selected
        ]
        task["execute_status"] = "stopping" if task.get("execute_stop_requested") else "running"
        task["execute_action"] = action
        task["execute_message"] = "正在服务器更新智赢产品" if action == "zying" else "正在更新泽顺数据和链接"
    if task.get("execute_stop_requested"):
        _finish_stopped_execute(task)
        return
    execution_id = uuid.uuid4().hex
    settings_cache = {}
    for row in records:
        row.update(_execution_id=execution_id, _execution_action=action, _execution_operator=owner)
        row["_history_context"] = _order_history_context(row, settings_cache)
        for key in ("_submitted_weight_g", "_submitted_dimensions_cm", "_dimensions_preserved"):
            row.pop(key, None)
    original_records = records
    if action == "zeshun":
        authorized_ids = task.get("authorized_ids", sorted({
            int(row["_token_id"]) for row in records if row.get("_token_id")
        }))
        records = _expand_same_title_records(
            records, authorized_ids, settings_cache=settings_cache,
            stop_requested=lambda: task.get("execute_stop_requested"),
        )
    erp_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    dim_groups: dict[tuple[int, str], list[dict[str, Any]]] = defaultdict(list)
    net_groups: dict[tuple[int, str], list[dict[str, Any]]] = defaultdict(list)
    for row in records:
        if row.get("_expansion_error") or row.get("_expansion_skip"):
            continue
        public = _public_record(row)
        if action == "zeshun" and not public["can_execute_zeshun"]:
            continue
        if action == "zying" and not public["can_execute_zying"]:
            continue
        if action == "combined" and not public["can_execute"]:
            continue
        product_id = str(row.get("product_id") or "").strip()
        weight = str(row.get("actual_weight_g") or "").strip()
        dimensions = str(row.get("actual_dimensions_cm") or "").strip()
        if (action != "zeshun" and not product_id) or not weight:
            row["execution_status"] = "未执行"
            row["execution_error"] = "缺少智赢产品编号或经官方核验的实际重量"
            _log_execution(row, "准备", "跳过", row["execution_error"])
            continue

        if action != "zeshun":
            erp_groups[product_id].append(row)
        if action == "zying":
            continue
        if row.get("_token_id") and row.get("_global_item_id") and row.get("marketplace_item_id"):
            token_id = int(row["_token_id"])
            dim_groups[(token_id, str(row["_global_item_id"]))].append(row)
            net_groups[(token_id, str(row["marketplace_item_id"]))].append(row)
        elif action != "zeshun":
            _log_execution(row, "美客多链接", "跳过", "该记录没有完整的美客多商品关联信息，仅处理智赢产品")

    latest_order_by_product = {
        product_id: _latest_order_row(grouped)
        for product_id, grouped in erp_groups.items()
    }

    if not erp_groups and not dim_groups:
        if action == "zeshun":
            _finish_same_title_records(original_records, records, stopped=bool(task.get("execute_stop_requested")))
        if task.get("execute_stop_requested"):
            _finish_stopped_execute(task)
            return
        with _lock:
            task["execute_status"] = "completed"
            task["execute_message"] = "没有可更新的链接，请查看订单执行日志"
            task["execute_total"] = 0
        return

    client_templates: dict[int, MercadoLibreClient] = {}
    errors: dict[int, str] = {}
    for token_id, _ in set(dim_groups) | set(net_groups):
        if token_id not in client_templates and token_id not in errors:
            try:
                token = bit_mysql.get_mercado_store_token(token_id)
                client_templates[token_id], _ = _client_and_token(dict(token or {}))
            except Exception as exc:
                errors[token_id] = str(exc)

    api_worker_local = threading.local()

    def api_client(token_id: int) -> MercadoLibreClient:
        clients = getattr(api_worker_local, "clients", None)
        if clients is None:
            clients = api_worker_local.clients = {}
        client = clients.get(token_id)
        if client is None:
            template = client_templates[token_id]
            client = MercadoLibreClient(template.access_token, timeout=template.timeout)
            clients[token_id] = client
        return client

    executed = 0
    total_operations = sum(len(rows) for rows in erp_groups.values())
    total_operations += sum(len(rows) for rows in dim_groups.values())
    total_operations += sum(len(rows) for rows in net_groups.values())
    with _lock:
        task["execute_total"] = total_operations
        task["execute_processed"] = 0
    progress_lock = threading.Lock()

    def advance_progress(count: int, message: str) -> None:
        nonlocal executed
        with progress_lock:
            executed += count
            with _lock:
                task["execute_processed"] = executed
                task["execute_message"] = f"{message} {executed}/{total_operations}"
    browser_row = {"current": None}

    def browser_logger(message, *args, **kwargs):
        current = browser_row.get("current")
        if current is not None:
            _log_execution(current, "智赢插件", str(kwargs.get("level") or "记录"), str(message))

    for product_id, grouped in erp_groups.items():
        if task.get("execute_stop_requested"):
            break
        row = latest_order_by_product[product_id]
        chosen_order = str(row.get("order_number") or "未知")
        chosen_time = str(row.get("time") or "时间未记录")
        for item in grouped:
            item["_zying_status"] = "执行中"
            item["execution_status"] = "执行中"
            item["execution_error"] = ""
            _log_execution(
                item, "智赢插件", "进行中",
                f"按产品编号 {product_id} 进入智赢产品页；该编号关联 {len(grouped)} 笔订单，采用最新订单 {chosen_order}（{chosen_time}）的数据：重量 {row['actual_weight_g']}g" + (f"，尺寸 {row['actual_dimensions_cm']}cm" if row.get("actual_dimensions_cm") else "；官方未提供尺寸，仅更新重量"),
            )
        try:
            if erp_browser_factory is None:
                raise RuntimeError("智赢插件自动化未配置")
            submitted_weight = package_update_weight_g(row)
            for item in grouped:
                item.update(_submitted_weight_g=submitted_weight,
                            _submitted_dimensions_cm=row.get("actual_dimensions_cm") or "")
            with erp_browser_factory(browser_logger) as erp_browser:
                browser_row["current"] = row
                erp_browser.update_package_by_product_id(
                    product_id, submitted_weight, row["actual_dimensions_cm"]
                )
            for item in grouped:
                item["_zying_status"] = "成功"
                item["execution_status"] = (
                    "智赢产品已更新"
                )
                _log_execution(
                    item, "智赢插件", "成功",
                    f"产品 {product_id} 提交重量 {submitted_weight}g" + (f"、尺寸 {row['actual_dimensions_cm']}cm" if row.get("actual_dimensions_cm") else "（官方未提供尺寸，沿用原值）") + f"、级别“重点”已保存并回读确认；采用最新订单 {chosen_order}",
                )
        except Exception as exc:
            for item in grouped:
                item["_zying_status"] = "失败"
                item["_package_status"] = "失败"
                item["execution_status"] = "失败"
                item["execution_error"] = f"智赢产品更新失败：{str(exc)[:250]}"
                _log_execution(item, "智赢插件", "失败", item["execution_error"])
        finally:
            browser_row["current"] = None
        advance_progress(len(grouped), "智赢产品更新")

    def update_dimensions_group(key, grouped):
        token_id, global_id = key
        eligible = (
            list(grouped)
            if action == "zeshun"
            else [row for row in grouped if row.get("_zying_status") == "成功"]
        )
        if not eligible:
            for item in grouped:
                if action == "zeshun" or (
                    item.get("_zying_status") != "成功" and item.get("_package_status") != "冲突"
                ):
                    _log_execution(item, "美客多链接", "跳过", "智赢产品未更新成功，按执行顺序跳过美客多链接更新")
            return
        if action != "zeshun" and len(eligible) != len(grouped):
            for item in grouped:
                if item.get("_zying_status") != "成功":
                    _log_execution(item, "美客多链接", "跳过", "智赢产品未更新成功，按执行顺序跳过美客多链接更新")
        try:
            if token_id in errors:
                raise RuntimeError("店铺授权不可用：" + errors[token_id])
            row = _latest_order_row(eligible)
            chosen_order = str(row.get("order_number") or "未知")
            chosen_time = str(row.get("time") or "时间未记录")
            submitted_weight = _marketplace_package_weight_g(package_update_weight_g(row))
            for item in eligible:
                item.update(_submitted_weight_g=submitted_weight,
                            _submitted_dimensions_cm=row.get("actual_dimensions_cm") or "",
                            _dimensions_preserved=not bool(row.get("actual_dimensions_cm")))
                _log_execution(
                    item, "美客多链接", "进行中",
                    f"正在更新 Global 商品 {global_id}；采用最新成功订单 {chosen_order}（{chosen_time}）的提交重量 {submitted_weight}g" + (f"、尺寸 {row['actual_dimensions_cm']}cm" if row.get("actual_dimensions_cm") else "；官方未提供尺寸，仅更新重量"),
                )
            if row.get("_automation_run_id"):
                from bit.automation_center import require_enabled
                require_enabled(row["_automation_organization"], "freight")
            _package_attributes(
                api_client(token_id), global_id, submitted_weight, row["actual_dimensions_cm"]
            )
            if row.get("_automation_run_id"):
                from bit.automation_center import backfill_local_package, journal
                for item in eligible:
                    result = backfill_local_package({**item, "actual_dimensions_cm": row.get("actual_dimensions_cm")}, submitted_weight)
                    journal(item["_automation_run_id"], {"stage": "泽顺同名产品回填", **result})
            for item in eligible:
                item["_package_status"] = "成功"
                item["_zeshun_status"] = "重量尺寸已更新" if row.get("actual_dimensions_cm") else "重量已更新"
                item["execution_status"] = "重量已更新，等待更新净收益" if not row.get("actual_dimensions_cm") else "重量尺寸已更新，等待更新净收益"
                _log_execution(item, "美客多链接", "成功", f"关联 Global 商品已更新，采用最新订单 {chosen_order} 的提交重量 {submitted_weight}g" + (f"和尺寸 {row['actual_dimensions_cm']}" if row.get("actual_dimensions_cm") else "；官方未提供尺寸，保留商品原尺寸"))
        except Exception as exc:
            for item in eligible:
                item["_package_status"] = "失败"
                item["_zeshun_status"] = "失败"
                item["execution_status"] = "部分完成"
                item["execution_error"] = f"美客多链接重量尺寸更新失败：{_update_error(exc)}"
                _log_execution(item, "美客多链接", "失败", item["execution_error"])

    def run_parallel_group_updates(groups, callback, stage_name, message):
        if task.get("execute_stop_requested"):
            return
        def run_unless_stopped(key, grouped):
            if not task.get("execute_stop_requested"):
                callback(key, grouped)
        items = list(groups.items())
        if not items or task.get("execute_stop_requested"):
            return
        workers = _configured_workers(
            len(items), UPDATE_MAX_WORKERS, UPDATE_MAX_WORKERS,
            "WEIGHT_DIMENSIONS_UPDATE_WORKERS",
        )
        # The automation worker owns a four-slot DB pool and keeps one slot for
        # its cross-process lease. Limit fan-out so parallel callbacks plus the
        # reserved authorization slot cannot exhaust that pool.
        if str(task.get("owner") or "").startswith("automation:"):
            workers = min(workers, 2)
        with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="wdr-update") as executor:
            future_groups = {}
            for key, grouped in items:
                if task.get("execute_stop_requested"):
                    break
                future_groups[executor.submit(run_unless_stopped, key, grouped)] = grouped
            for future in as_completed(future_groups):
                grouped = future_groups[future]
                try:
                    future.result()
                except Exception as exc:
                    for row in grouped:
                        row["execution_status"] = "部分完成"
                        row["execution_error"] = f"{stage_name}更新失败：{str(exc)[:250]}"
                        _log_execution(row, stage_name, "失败", row["execution_error"])
                advance_progress(len(grouped), message)

    # These independent marketplace groups use separate HTTP sessions per worker.
    # Wait for all dimension writes before re-submitting any current net proceeds.
    run_parallel_group_updates(
        dim_groups, update_dimensions_group, "美客多链接", "智赢与美客多商品重量尺寸更新",
    )

    # All package attributes are submitted before the second pass changes net proceeds.
    def update_net_group(key, grouped):
        token_id, marketplace_id = key
        eligible = [row for row in grouped if row.get("_package_status") == "成功"]
        if not eligible:
            for item in grouped:
                if item.get("_package_status") != "冲突":
                    _log_execution(item, "净收益", "跳过", "美客多链接重量尺寸未更新成功，按执行顺序跳过净收益重提")
            return
        if token_id in errors:
            for row in eligible:
                row["_net_status"] = "失败"
                row["_zeshun_status"] = "失败"
                row["execution_status"] = "部分完成"
                row["execution_error"] = "重量尺寸已提交，但店铺授权不可用：" + errors[token_id]
                _log_execution(row, "净收益", "失败", row["execution_error"])
            return
        try:
            first = eligible[0]
            _log_execution(first, "净收益", "进行中", "尺寸更新后读取商品当前 USD 净收益")
            client = api_client(token_id)
            current_item = client.get_marketplace_item(
                marketplace_id, attributes=("id", "net_proceeds")
            )
            target = _extract_net_proceeds(current_item)
            if target is None:
                raise ValueError("商品当前没有可读取的 USD 净收益")
            for row in eligible:
                row["current_net_proceeds_usd"] = str(target)
                row["_net_proceeds_usd"] = target
                _log_execution(row, "净收益", "进行中", f"读取当前净收益 USD {target}，正在重新提交")
            if row.get("_automation_run_id"):
                from bit.automation_center import require_enabled
                require_enabled(row["_automation_organization"], "freight")
            client.update_global_item(marketplace_id, {"net_proceeds": float(target)})
            for row in eligible:
                row["_net_status"] = "成功"
                row["_zeshun_status"] = "成功"
                row["execution_status"] = "完成"
                row["execution_error"] = ""
                _log_execution(row, "净收益", "成功", f"已重新提交当前净收益 USD {target}")
        except Exception as exc:
            for row in eligible:
                row["_net_status"] = "失败"
                row["_zeshun_status"] = "失败"
                row["execution_status"] = "部分完成"
                row["execution_error"] = f"净收益更新失败：{str(exc)[:250]}"
                _log_execution(row, "净收益", "失败", row["execution_error"])

    run_parallel_group_updates(
        net_groups, update_net_group, "净收益", "尺寸与净收益更新",
    )
    if action == "zeshun":
        _finish_same_title_records(original_records, records, stopped=bool(task.get("execute_stop_requested")))
        records = original_records
    if task.get("execute_stop_requested"):
        _finish_stopped_execute(task)
        return
    with _lock:
        task["execute_status"] = "completed"
        if action == "zeshun":
            completed_count = sum(row.get("_zeshun_status") == "成功" for row in records)
            task["execute_message"] = (f"泽顺数据和链接更新结束，共 {len(records)} 条订单，"
                                       f"全部成功 {completed_count} 条，"
                                       f"未全部完成 {len(records) - completed_count} 条；请查看逐条日志")
        else:
            completed_count = sum(row.get("_zying_status") == "成功" for row in records)
            failed_rows = [row for row in records if row.get("_zying_status") == "失败"]
            task["execute_message"] = (
                f"智赢产品更新结束，共 {len(records)} 条订单，成功 {completed_count} 条，"
                f"失败 {len(failed_rows)} 条"
            )
            if failed_rows:
                task["execute_message"] += "；首个失败原因：" + str(failed_rows[0].get("execution_error") or "请查看逐条日志")


def _run_execute_guarded(task_id, owner, action, selected_order_numbers, erp_browser_factory):
    """Ensure an unexpected worker exception cannot leave rows marked executing."""
    try:
        _run_execute(task_id, owner, action, selected_order_numbers, erp_browser_factory)
    except Exception as exc:
        _logger.exception("重量尺寸执行任务 %s 异常结束（action=%s）", task_id, action)
        task = _tasks.get(task_id)
        if not task:
            return
        selected = {str(value or "").strip() for value in selected_order_numbers or () if str(value or "").strip()}
        reason = f"重量尺寸更新异常结束：{str(exc)[:400]}"
        affected = []
        with _lock:
            for row in task.get("records") or ():
                if selected and str(row.get("order_number") or "").strip() not in selected:
                    continue
                zying_done = row.get("_zying_status") in {"成功", "失败", "跳过"}
                zeshun_done = row.get("_zeshun_status") in {"成功", "失败", "跳过"}
                if ((action == "zying" and zying_done)
                        or (action == "zeshun" and zeshun_done)
                        or (action == "combined" and zying_done and zeshun_done)):
                    continue
                row["execution_status"] = "失败"
                row["execution_error"] = reason
                if action in {"zying", "combined"} and row.get("_zying_status") != "成功":
                    row["_zying_status"] = "失败"
                if action in {"zeshun", "combined"} and row.get("_zeshun_status") != "成功":
                    row["_zeshun_status"] = "失败"
                affected.append(row)
            task["execute_status"] = "completed"
            task["execute_message"] = f"{reason}；未完成记录已标记失败"
        for row in affected:
            _log_execution(row, "任务", "失败", reason)
    finally:
        with _lock:
            if _execution_threads.get(task_id) is threading.current_thread():
                _execution_threads.pop(task_id, None)


def start_execute(
    task_id: str,
    owner: str,
    *,
    action="zeshun",
    selected_order_numbers=None,
    select_all_matching=False,
    excluded_order_numbers=None,
    erp_browser_factory=None,
    erp_browser_preflight=None,
    agent_dispatch=None,
) -> dict[str, Any]:
    task = _get_owned_task(task_id, owner)
    action = str(action or "zeshun").strip().lower()
    if action not in {"zeshun", "zying", "combined"}:
        raise ValueError("更新动作无效")
    use_agent = action == "zying" and agent_dispatch is not None
    selected = {
        str(value or "").strip()
        for value in selected_order_numbers or ()
        if str(value or "").strip()
    }
    excluded = {
        str(value or "").strip()
        for value in excluded_order_numbers or ()
        if str(value or "").strip()
    } if select_all_matching else set()
    saved_selection = None
    if task.get("database_view"):
        saved_selection = bit_db_api.list_weight_dimensions_records(
            None if select_all_matching else (sorted(selected) if selected else None),
            filters=task["filters"],
        )
    with _lock:
        if task.get("status") != "ready":
            raise ValueError("订单查询完成后才能执行更新")
        if task.get("execute_status") in {"queued", "running", "stopping"}:
            raise ValueError("当前任务正在执行")
        if _executions_shutting_down and not use_agent:
            raise ValueError("服务正在关闭，请重新连接后执行重量尺寸更新")
        if action in {"zeshun", "zying"} and not selected and not select_all_matching:
            raise ValueError("请至少勾选一条需要更新的订单")
        candidates = [
            row for row in (saved_selection if saved_selection is not None else task["records"])
            if (select_all_matching or not selected or str(row.get("order_number") or "").strip() in selected)
            and str(row.get("order_number") or "").strip() not in excluded
        ]
        if action == "zeshun":
            ready_rows = [row for row in candidates if _public_record(row)["can_execute_zeshun"]]
        elif action == "zying":
            ready_rows = [row for row in candidates if _public_record(row)["can_execute_zying"]]
        else:
            ready_rows = [row for row in candidates if _public_record(row)["can_execute"]]
        ready = len(ready_rows)
        if not ready:
            if action == "combined":
                all_already_succeeded = bool(candidates) and all(
                    row.get("_zeshun_status") == "成功" and row.get("_zying_status") == "成功"
                    for row in candidates
                )
            else:
                status_key = "_zeshun_status" if action == "zeshun" else "_zying_status"
                all_already_succeeded = bool(candidates) and all(
                    row.get(status_key) == "成功" for row in candidates
                )
            if all_already_succeeded:
                task["execute_status"] = "completed"
                task["execute_action"] = action
                task["execute_message"] = f"所选 {len(candidates)} 条记录已成功，已跳过，无需重复执行"
                task["execute_total"] = 0
                task["execute_processed"] = 0
                task["execute_order_numbers"] = []
                task["agent_job_id"] = ""
                return {
                    "status": "completed", "ready_count": 0,
                    "skipped_success_count": len(candidates),
                }
            raise ValueError(
                "所选记录没有具备可执行数据；智赢更新需要产品id和实际重量，"
                "泽顺更新需要美客多链接关联和实际重量；尺寸有官方核验值时一并更新"
            )
        selected = {str(row.get("order_number") or "").strip() for row in ready_rows}

    # Check the execution computer's visible browser before queuing any writes.
    # Opening Edge and waiting for login must not hold the global task lock.
    if action in {"zying", "combined"} and not use_agent and erp_browser_preflight is not None:
        erp_browser_preflight()

    with _lock:
        if _tasks.get(str(task_id)) is not task:
            raise KeyError("查询任务不存在或已过期")
        if task.get("status") != "ready":
            raise ValueError("订单查询完成后才能执行更新")
        if task.get("execute_status") in {"queued", "running", "stopping"}:
            raise ValueError("当前任务正在执行")
        if _executions_shutting_down and not use_agent:
            raise ValueError("服务正在关闭，请重新连接后执行重量尺寸更新")
        if saved_selection is not None:
            task["records"] = saved_selection
        task["execute_status"] = "queued"
        task["execute_stop_requested"] = False
        task.pop("agent_cancel", None)
        task["execute_action"] = action
        task["execute_message"] = f"已排队，准备执行 {ready} 条记录"
        task["execute_total"] = ready
        task["execute_processed"] = 0
        task["agent_job_id"] = ""
        task["execute_order_numbers"] = sorted(selected)
        task["agent_applied_status"] = ""
        task["agent_applied_result"] = False
        if not use_agent:
            thread = threading.Thread(
                target=_run_execute_guarded,
                args=(task_id, owner, action, selected, erp_browser_factory),
                name=f"weight-dimensions-update-{task_id[:8]}", daemon=True,
            )
            _execution_threads[task_id] = thread
            try:
                thread.start()
            except Exception:
                _execution_threads.pop(task_id, None)
                task["execute_status"] = "idle"
                task["execute_message"] = "重量尺寸执行线程启动失败"
                raise

    if use_agent:
        execution_id = uuid.uuid4().hex
        settings_cache = {}
        for row in ready_rows:
            row.update(_execution_id=execution_id, _execution_action=action,
                       _execution_operator=owner)
            row["_history_context"] = _order_history_context(row, settings_cache)
        payload_rows = [
            {
                "order_number": row.get("order_number"),
                "product_id": row.get("product_id"),
                "actual_weight_g": row.get("actual_weight_g"),
                "declared_weight_g": row.get("declared_weight_g"),
                "actual_dimensions_cm": row.get("actual_dimensions_cm"),
            }
            for row in ready_rows
        ]
        try:
            job_id = str(agent_dispatch(payload_rows) or "").strip()
            if not job_id:
                raise ValueError("Agent 未返回任务编号")
        except Exception:
            with _lock:
                task["execute_status"] = "idle"
                task["execute_message"] = "智赢 Agent 任务提交失败"
            raise
        with _lock:
            task["agent_job_id"] = job_id
            task["execute_status"] = "stopping" if task.get("execute_stop_requested") else "queued"
            task["execute_message"] = f"已提交智赢 Agent，准备逐个更新 {ready} 个产品"
            for row in ready_rows:
                row["_zying_status"] = "排队中"
                row["execution_status"] = "排队中"
                row["execution_error"] = ""
        for row in ready_rows:
            _log_execution(row, "智赢 Agent", "排队中", "任务已提交本机 Agent 队列，尚未开始更新")
        if task.get("execute_stop_requested") and task.get("agent_cancel"):
            task["agent_cancel"](job_id)
        return {"status": task["execute_status"], "ready_count": ready, "agent_job_id": job_id}

    return {"status": "queued", "ready_count": ready}
