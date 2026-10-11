"""Parse Zying exports without creating orders or guessing pack cost allocation."""
import io
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20000
HEADERS = {
    "order_id": ("平台订单号", "美客多订单号", "原始订单号", "订单ID", "订单id", "order_no", "订单号", "订单编号", "销售单号"),
    "purchase_order": ("找货订单号", "采购订单号", "采购单号", "purchase_order"),
    "purchase_tracking": ("采购物流号", "采购物流单号", "采购追踪号", "采购追踪", "物流追踪号", "物流单号", "物流号", "运单号", "purchase_tracking"),
    "purchase_cost": ("采购成本", "订单成本", "人民币成本", "成本（人民币）", "成本(RMB)", "成本", "order_cost", "purchase_cost"),
}


def header(value):
    name = re.sub(r"\s+", "", str(value or "")).replace("（", "(").replace("）", ")").casefold()
    return re.sub(r"\((?:人民币|rmb|cny|元|¥|￥)\)$", "", name)


def identifier(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        raise ValueError("订单号或物流号格式异常")
    if isinstance(value, (int, float)):
        if not Decimal(str(value)).is_finite() or value != int(value):
            raise ValueError("订单号或物流号必须为完整编号")
        if abs(value) >= 10 ** 15:
            raise ValueError("长编号使用了数值单元格，可能丢失精度；请在源表中将编号设为文本后重新导出")
        return str(int(value))
    text = str(value).strip()
    if text.startswith("="):
        raise ValueError("编号不能使用公式，请转换为文本")
    return text


def money(value):
    text = re.sub(r"^(?:人民币|RMB|CNY|[¥￥])\s*", "", str(value).strip(), flags=re.I)
    try:
        if "," in text and not re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", text):
            raise InvalidOperation
        amount = Decimal(text.replace(",", ""))
        if not amount.is_finite() or amount < 0 or amount >= Decimal("10000000000"):
            raise InvalidOperation
        return str(amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    except InvalidOperation as exc:
        raise ValueError("人民币采购成本无效") from exc


def purchase_snapshot(buys):
    """Replace purchase fields with the complete, numbered purchase snapshot."""
    amounts, numbers, tracking = {}, [], []
    for buy in buys:
        if not isinstance(buy, dict):
            raise ValueError("智赢采购数据格式异常")
        number = identifier(buy.get("no"))
        if not number:
            continue
        parts = list(dict.fromkeys(part for part in re.split(r"[,，;；\s]+", number) if part))
        if not parts:
            continue
        key = tuple(sorted(parts))
        currency = str(buy.get("cur") or "CNY").strip().upper()
        if currency not in {"CNY", "RMB", "人民币", "元", "¥", "￥"}:
            raise ValueError("采购单金额不是人民币，无法汇总成本")
        if buy.get("amount") in (None, ""):
            raise ValueError("采购单金额缺失，请更新泽顺插件并重新同步；Excel 请填写采购成本")
        amount = Decimal(money(buy["amount"]))
        if key in amounts and amounts[key] != amount:
            raise ValueError("同一采购单存在不同成本，请核对后重新同步")
        if key not in amounts and any(set(key).intersection(old) for old in amounts):
            raise ValueError("采购单组合重叠，无法确定成本归属")
        amounts[key] = amount
        numbers.extend(parts)
        tracking.extend(part for part in re.split(r"[,，;；\s]+", identifier(buy.get("trace"))) if part)
    changes = {"purchase_order": ",".join(dict.fromkeys(numbers)),
               "purchase_tracking": ",".join(dict.fromkeys(tracking)),
               "purchase_cost": money(sum(amounts.values(), Decimal(0)))}
    if any(len(changes[key]) > 255 for key in ("purchase_order", "purchase_tracking")):
        raise ValueError("采购单号或物流号超过 255 个字符")
    return changes


def read_excel(content, filename):
    if Path(filename).suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("请上传 .xlsx 或 .xlsm 文件；旧版 .xls 请先另存为 .xlsx")
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("Excel 文件不能为空且不能超过 10 MB")
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Excel 解压后过大，请分批上传")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError("无法读取 Excel，请确认文件完整且未加密") from exc
    records, count, sheets = {}, 0, 0
    try:
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            columns = None
            for header_number in range(1, 21):
                values = next(iterator, None)
                if values is None:
                    break
                names = [header(value) for value in values]
                resolved = {key: next((names.index(header(alias)) for alias in aliases if header(alias) in names), None)
                            for key, aliases in HEADERS.items()}
                if resolved["order_id"] is not None and any(resolved[key] is not None for key in HEADERS if key != "order_id"):
                    columns = resolved
                    break
            if columns is None:
                continue
            sheets += 1
            for row_number, values in enumerate(iterator, header_number + 1):
                if not any(value not in (None, "") for value in values):
                    continue
                count += 1
                if count > MAX_ROWS:
                    raise ValueError("单次最多上传 20,000 行，请拆分文件")
                location = f"{sheet.title} 第 {row_number} 行"
                raw = {key: values[index] if index is not None and index < len(values) else None for key, index in columns.items()}
                try:
                    order_id = identifier(raw["order_id"])
                except ValueError as exc:
                    order_id = ""
                    identity_error = str(exc)
                else:
                    identity_error = "缺少平台订单号" if not order_id else ""
                record = records.setdefault(order_id or f"__row_{count}", {
                    "order_id": order_id, "changes": {}, "locations": [], "error": "", "_buys": [],
                })
                record["locations"].append(location)
                try:
                    if identity_error:
                        raise ValueError(identity_error)
                    if columns["purchase_order"] is None or columns["purchase_cost"] is None:
                        raise ValueError("请同时提供采购单号和采购成本列，成本按每笔采购单填写")
                    record["_buys"].append({"no": raw["purchase_order"],
                        "trace": raw["purchase_tracking"], "amount": raw["purchase_cost"]})
                    record["changes"] = purchase_snapshot(record["_buys"])
                except ValueError as exc:
                    record["error"] = f"{location}：{exc}"
    finally:
        workbook.close()
    if not sheets:
        raise ValueError("未识别到订单表：需要平台订单号（或订单号）以及采购成本、采购单号、物流号中的至少一列")
    if not records:
        raise ValueError("Excel 中没有订单数据")
    for record in records.values():
        record.pop("_buys", None)
    return list(records.values())
