from pathlib import Path

import openpyxl


MAPPING = Path(r"E:\xwechat_files\z1013459852_e03a\temp\RWTemp\2026-09\12f3855f62071e413148208cd1d773e4\店铺流水.xlsx")
SOURCE_DIR = Path(r"E:\360MoveData\Users\Admin\Documents\20260911")


def show_book(path: Path, max_rows: int = 12, max_cols: int = 25) -> None:
    print(f"\n=== {path} ===")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    print("sheets:", wb.sheetnames)
    for ws in wb.worksheets:
        print(f"-- {ws.title}: rows={ws.max_row}, cols={ws.max_column}")
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows), values_only=True):
            print(tuple(row[:max_cols]))


show_book(MAPPING, max_rows=20)
files = sorted((p for p in SOURCE_DIR.glob("*.xlsx") if not p.name.startswith("~$")), key=lambda p: p.name)
for path in files[:2] + files[len(files) // 2:len(files) // 2 + 1] + files[-1:]:
    show_book(path, max_rows=8, max_cols=40)
