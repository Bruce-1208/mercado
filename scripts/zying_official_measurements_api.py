"""通过美客多官方运费补差 API 导出智赢订单的实测重量和尺寸。

智赢导出文件只提供订单号映射；其中毛重、尺寸列即使为空也不会使用。
不连接或控制智赢桌面界面，不修改订单和商品。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from bit import bit_mysql, mercado_tokens  # noqa: E402
from mercado_api.client import MercadoLibreClient  # noqa: E402


def _id(value, minimum=7, maximum=22):
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text if re.fullmatch(rf"\d{{{minimum},{maximum}}}", text) else ""


def load_order_pairs(path, limit):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.values
        header = next(rows, None)
        if header is None:
            raise ValueError("订单导出文件为空")
        names = [str(value or "").strip().casefold() for value in header]
        if names.count("id") != 1 or names.count("编号") != 1:
            raise ValueError("智赢导出文件必须有 id 和 编号 两列")
        internal_col, platform_col = names.index("id"), names.index("编号")
        pairs, seen = [], set()
        for row in rows:
            internal = _id(row[internal_col] if internal_col < len(row) else None, 7, 12)
            platform = _id(row[platform_col] if platform_col < len(row) else None, 12, 22)
            if not internal or not platform or internal in seen:
                continue
            seen.add(internal)
            pairs.append((internal, platform))
            if len(pairs) == limit:
                break
        if not pairs:
            raise ValueError("导出文件中没有同时带智赢 id 和平台编号的订单")
        return pairs
    finally:
        workbook.close()


def local_token_map(platform_ids):
    """从已同步订单找对应的店铺授权；只读取映射，不读取或打印密钥。"""
    result = {}
    connection = bit_mysql.pymysql.connect(**bit_mysql.config)
    try:
        with connection.cursor() as cursor:
            for start in range(0, len(platform_ids), 200):
                batch = platform_ids[start:start + 200]
                placeholders = ",".join(["%s"] * len(batch))
                cursor.execute(
                    "SELECT order_id, token_id FROM mercado_synced_orders "
                    f"WHERE order_id IN ({placeholders})",
                    batch,
                )
                for row in cursor.fetchall():
                    order_id = str(row["order_id"])
                    token_id = int(row["token_id"])
                    if order_id in result and result[order_id] != token_id:
                        raise RuntimeError(f"平台订单 {order_id} 对应多个店铺授权")
                    result[order_id] = token_id
    finally:
        connection.close()
    return result


def _refresh_if_needed(token):
    expiry = token.get("expires_at")
    if isinstance(expiry, str):
        try:
            expiry = datetime.fromisoformat(expiry)
        except ValueError:
            expiry = None
    if expiry and expiry.tzinfo is not None:
        expiry = expiry.astimezone().replace(tzinfo=None)
    if expiry and expiry <= datetime.now() + timedelta(minutes=5):
        mercado_tokens.refresh_and_save(
            int(token["id"]),
            get_token=bit_mysql.get_mercado_store_token,
            update_token=bit_mysql.update_mercado_store_token,
            record_error=bit_mysql.record_mercado_store_token_error,
        )
        token = bit_mysql.get_mercado_store_token(int(token["id"]))
    return token


def _number(value):
    try:
        number = Decimal(str(value))
    except (InvalidOperation, TypeError) as exc:
        raise ValueError(f"实测值无效：{value!r}") from exc
    if not number.is_finite() or number <= 0:
        raise ValueError(f"实测值无效：{value!r}")
    return number


def validated_measurement(data):
    """只取 validated.net 与 validated.dimensions，绝不回退到 declared/billable。"""
    validated = ((data or {}).get("package") or {}).get("validated") or {}
    weight = validated.get("weight") or validated.get("weights") or {}
    dimensions = validated.get("dimensions") or {}
    if str(weight.get("unit") or "").lower() != "g":
        raise ValueError("API 没有以克返回实测重量")
    if str(dimensions.get("unit") or "").lower() != "cm":
        raise ValueError("API 没有以厘米返回实测尺寸")
    net = _number(weight.get("net"))
    lengths = [_number(dimensions.get(key)) for key in ("length", "width", "height")]
    plain = lambda value: format(value.normalize(), "f")
    return plain(net), "x".join(plain(v) for v in lengths)


def fetch_one(client, platform_order_id):
    # 智赢导出的“编号”是包裹号；普通订单接口会对此类编号返回 404。
    pack = client.request("GET", f"/marketplace/orders/pack/{platform_order_id}")
    if str(pack.get("id")) != platform_order_id:
        raise ValueError("API 返回的包裹号不匹配")
    shipment_id = _id((pack.get("shipment") or {}).get("id"), 5, 22)
    if not shipment_id:
        raise ValueError("订单没有可查询的运单号")
    compensation = client.request(
        "GET", f"/marketplace/shipments/{shipment_id}/compensation_costs",
        params={"weight_unit": "g", "dimensions_unit": "cm"},
    )
    return validated_measurement(compensation)


def write_excel(records, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "官方实测重量尺寸"
    sheet.append(["订单号", "重量（克）", "尺寸（厘米）"])
    for record in records:
        weight = Decimal(record["weight_g"])
        sheet.append([record["platform_order_id"],
                      int(weight) if weight == int(weight) else float(weight),
                      record["dimensions_cm"]])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 23
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet["A"][1:]:
        cell.number_format = "@"
    for cell in sheet["B"][1:]:
        cell.number_format = "0.##"
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, help="智赢自带订单导出 .xlsx")
    parser.add_argument("--list-stores", action="store_true", help="只列出已配置店铺的 ID 和名称")
    parser.add_argument("--max-count", type=int, default=100)
    parser.add_argument("--token-id", type=int,
                        help="指定一个已授权的美客多店铺；不指定则用本地已同步订单自动匹配")
    parser.add_argument("--output", type=Path,
                        default=Path(f"output/zying_official_api_{datetime.now():%Y%m%d_%H%M%S}.xlsx"))
    args = parser.parse_args()
    if args.list_stores:
        for store in bit_mysql.list_mercado_store_tokens()["rows"]:
            if store.get("enabled"):
                print(f'{store["id"]}\t{store.get("display_name") or store.get("nickname") or "未命名"}')
        return
    if args.input is None:
        parser.error("请提供 --input 智赢订单导出文件")
    if args.max_count < 1:
        parser.error("max-count 必须大于0")
    if args.output.suffix.lower() != ".xlsx":
        parser.error("output 必须以 .xlsx 结尾")
    pairs = load_order_pairs(args.input, args.max_count)
    token_map = {} if args.token_id else local_token_map([p for _, p in pairs])
    clients = {}
    log_path = args.output.with_suffix(".jsonl")
    log_path.parent.mkdir(parents=True, exist_ok=True)
    log_path.write_text("", encoding="utf-8")

    def log(status, **fields):
        event = {"time": datetime.now().isoformat(), "status": status, **fields}
        with log_path.open("a", encoding="utf-8") as stream:
            stream.write(json.dumps(event, ensure_ascii=False) + "\n")
        print(json.dumps(event, ensure_ascii=False), flush=True)

    records = []
    for internal_id, platform_id in pairs:
        token_id = args.token_id or token_map.get(platform_id)
        if not token_id:
            log("skipped", order_id=internal_id, reason="找不到对应店铺授权；可指定 --token-id")
            continue
        try:
            if token_id not in clients:
                token = bit_mysql.get_mercado_store_token(int(token_id))
                if not token or not token.get("access_token"):
                    raise ValueError("店铺未授权或授权已停用")
                token = _refresh_if_needed(token)
                clients[token_id] = MercadoLibreClient(str(token["access_token"]))
            weight_g, dimensions_cm = fetch_one(clients[token_id], platform_id)
        except Exception as exc:
            # 403/404、没有补差或没有实测值都不使用产品预设尺寸填补。
            log("skipped", order_id=internal_id, platform_order_id=platform_id,
                reason=str(exc)[:300])
            continue
        record = {"platform_order_id": platform_id, "weight_g": weight_g,
                  "dimensions_cm": dimensions_cm}
        records.append(record)
        log("read", **record)
    write_excel(records, args.output)
    log("finished", requested=len(pairs), exported=len(records), excel=str(args.output))


if __name__ == "__main__":
    main()
