from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re

import openpyxl


MAPPING = Path(r"E:\xwechat_files\z1013459852_e03a\temp\RWTemp\2026-09\12f3855f62071e413148208cd1d773e4\店铺流水.xlsx")
SOURCE_DIR = Path(r"E:\360MoveData\Users\Admin\Documents\20260911")


def page_no(path: Path) -> int:
    m = re.search(r"第(\d+)页", path.name)
    return int(m.group(1)) if m else 0


def canonical_store(value):
    if value is None:
        return None
    s = str(value).strip()
    s = re.sub(r"\([^)]*\)\s*$", "", s).strip()
    return s.casefold()


mapping_wb = openpyxl.load_workbook(MAPPING, read_only=True, data_only=True)
mapping_ws = mapping_wb["Sheet1"]
mapping = {}
duplicate_mapping = defaultdict(set)
for row in mapping_ws.iter_rows(min_row=2, values_only=True):
    store, company, legal_person, active_sites = row[:4]
    if not store or not company:
        continue
    key = canonical_store(store)
    duplicate_mapping[key].add(company)
    mapping[key] = (str(store).strip(), str(company).strip(), legal_person, active_sites)

files = sorted((p for p in SOURCE_DIR.glob("*.xlsx") if not p.name.startswith("~$")), key=page_no)
all_ids = []
store_counts = Counter()
store_amounts = defaultdict(Decimal)
page_info = []
bad_amount = []
blank_rows = 0

for path in files:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        order_id, order_no, salesperson, store_raw, amount_raw = r[:5]
        if all(v is None for v in r[:5]):
            blank_rows += 1
            continue
        if order_id is None:
            continue
        try:
            amount = Decimal(str(amount_raw).replace("$", "").replace(",", "").strip())
        except (InvalidOperation, AttributeError):
            bad_amount.append((path.name, order_id, amount_raw))
            continue
        key = canonical_store(store_raw)
        rows.append((int(order_id), str(order_no), salesperson, store_raw, amount))
        all_ids.append(int(order_id))
        store_counts[key] += 1
        store_amounts[key] += amount
    page_info.append((page_no(path), len(rows), rows[0] if rows else None, rows[-1] if rows else None, wb.properties.created, wb.properties.modified))

id_counts = Counter(all_ids)
unmatched = [(k, store_counts[k], store_amounts[k]) for k in store_counts if k not in mapping]
matched = [(k, store_counts[k], store_amounts[k], mapping[k]) for k in store_counts if k in mapping]

print(f"mapping rows={len(mapping)} duplicate keys={sum(1 for v in duplicate_mapping.values() if len(v)>1)}")
for key, companies in duplicate_mapping.items():
    if len(companies) > 1:
        print("DUPLICATE MAPPING", key, sorted(companies))
print(f"files={len(files)} records={len(all_ids)} unique_ids={len(id_counts)} duplicate_id_count={sum(1 for v in id_counts.values() if v>1)} blank_rows={blank_rows} bad_amounts={len(bad_amount)}")
print(f"unique_stores={len(store_counts)} matched_stores={len(matched)} unmatched_stores={len(unmatched)}")
print("\nBAD AMOUNTS")
for item in bad_amount[:120]:
    print(item)
print("\nPAGE INFO")
for item in page_info:
    print(item)
print("\nUNMATCHED BY AMOUNT")
for key, count, amount in sorted(unmatched, key=lambda x: x[2], reverse=True):
    print(key, count, amount)
print("\nMATCHED BY AMOUNT")
for key, count, amount, m in sorted(matched, key=lambda x: x[2], reverse=True):
    print(m[0], m[1], count, amount)
