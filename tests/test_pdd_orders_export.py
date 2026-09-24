import importlib.util
import json
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import Mock, patch

spec = importlib.util.spec_from_file_location('pdd_export', Path(__file__).resolve().parents[1] / 'scripts/pdd_orders.py')
pdd = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = pdd
spec.loader.exec_module(pdd)


class PddExportTests(unittest.TestCase):
    def test_exit_zero_waits_for_available_browser(self):
        process = Mock()
        process.poll.return_value = 0
        with patch.object(pdd, 'probe_cdp', side_effect=[None, 'http://127.0.0.1:1234']), patch.object(pdd.time, 'sleep'):
            self.assertEqual(pdd.wait_for_cdp(process, 1234), 'http://127.0.0.1:1234')

    def test_reuse_old_profile_without_launching_another_process(self):
        with patch.object(pdd, 'active_profile_endpoint', return_value=None), \
             patch.object(pdd, 'existing_profile_ports', return_value=(True, [1234])), \
             patch.object(pdd, 'probe_cdp', return_value='http://127.0.0.1:1234'), \
             patch.object(pdd.subprocess, 'Popen') as launch:
            process, endpoint = pdd.launch_edge_for_collection(Path('edge.exe'), Path('profile'))
            self.assertIsNone(process)
            self.assertEqual(endpoint, 'http://127.0.0.1:1234')
            launch.assert_not_called()

    def test_plain_existing_window_has_actionable_error(self):
        with patch.object(pdd, 'active_profile_endpoint', return_value=None), \
             patch.object(pdd, 'existing_profile_ports', return_value=(True, [])), \
             patch.object(pdd.subprocess, 'Popen') as launch:
            with self.assertRaisesRegex(RuntimeError, '正常关闭该专用窗口'):
                pdd.launch_edge_for_collection(Path('edge.exe'), Path('profile'))
            launch.assert_not_called()

    def test_original_error_and_json_survive_excel_lock(self):
        capture = pdd.Capture()
        capture.merge(pdd.Order('saved-order'))
        checkpoint = Mock()
        checkpoint.flush.side_effect = PermissionError('workbook locked')
        with tempfile.TemporaryDirectory() as directory, patch('builtins.print'):
            path = Path(directory)
            pdd.save_failure(capture, checkpoint, path, 'failed', RuntimeError('original browser failure'))
            self.assertEqual(json.loads((path / 'failed.json').read_text(encoding='utf-8'))[0]['order_id'], 'saved-order')
            self.assertIn('original browser failure', (path / 'failed.error.log').read_text(encoding='utf-8'))

    def test_fast_scroll_keeps_slow_loading_boundary(self):
        scroll = {'before': 0, 'after': 600, 'height': 10000, 'viewport': 800}
        self.assertEqual(pdd.scroll_wait_ms(scroll, False, 6), 500)
        self.assertEqual(pdd.scroll_wait_ms(scroll, True, 6), 6000)
        self.assertEqual(pdd.scroll_wait_ms(dict(scroll, after=9000), False, 6), 6000)
        self.assertEqual(pdd.scroll_wait_ms(dict(scroll, after=0), False, 6), 6000)

    def test_checkpoint_saves_json_each_round_and_excel_periodically(self):
        capture = pdd.Capture()
        capture.merge(pdd.Order('a', product='before'))
        with tempfile.TemporaryDirectory() as directory, patch.object(pdd.time, 'monotonic', return_value=0) as clock:
            folder = Path(directory)
            output = pdd.OutputCheckpoint(capture, folder, date(2026, 1, 1), date(2026, 9, 23), 'test')
            with patch.object(pdd, 'write_excel', return_value=folder / 'test.xlsx') as excel:
                output()
                self.assertEqual(excel.call_count, 1)
                capture.orders['a'].product = 'after'
                clock.return_value = 10
                output()
                self.assertEqual(excel.call_count, 1)
                self.assertEqual(json.loads((folder / 'test.json').read_text(encoding='utf-8'))[0]['product'], 'after')
                clock.return_value = 61
                output()
                self.assertEqual(excel.call_count, 2)
                pdd.checkpoint_now(output)
                self.assertEqual(excel.call_count, 3)

    def make_scrolling_page(self, last_page=245):
        page = Mock()
        page.url = pdd.ORDERS_URL
        page.locator.return_value.inner_text.return_value = '全部订单'
        state = {'position': 0}
        def evaluate(script):
            pos = state['position']
            if script == pdd.EXTRACT_CARDS_JS:
                return [{'text': f'订单号：260101-{pos:09d} 下单时间：2026-01-02'}]
            if script == pdd.LOAD_MORE_JS:
                return False
            if script == pdd.SCROLL_ORDERS_JS:
                state['position'] = min(pos + 1, last_page)
                return {'before': pos * 600, 'after': state['position'] * 600,
                        'height': 200000, 'viewport': 800, 'target': 'inner'}
            raise AssertionError(script)
        page.evaluate.side_effect = evaluate
        return page

    def test_unlimited_continues_past_240_rounds_and_pauses_on_stall(self):
        capture = pdd.Capture()
        with patch('builtins.input', return_value='q') as prompt, patch('builtins.print'):
            result = pdd.collect_orders(self.make_scrolling_page(), date(2026, 1, 1), date(2026, 9, 23),
                                        0, capture, lambda: None, interval=0)
        self.assertFalse(result)
        self.assertEqual(len(capture.orders), 246)
        prompt.assert_called_once()
        self.assertIn('未确认到底', capture.reason)

    def test_explicit_limit_reads_final_loaded_dom_and_marks_partial(self):
        capture = pdd.Capture()
        checkpoint = Mock()
        with patch('builtins.print'):
            result = pdd.collect_orders(self.make_scrolling_page(), date(2026, 1, 1), date(2026, 9, 23),
                                        2, capture, checkpoint, interval=0)
        self.assertFalse(result)
        self.assertEqual(len(capture.orders), 3)
        self.assertIn('列表未完成', capture.reason)
        self.assertIn('2026-01-02', capture.reason)
        self.assertTrue(checkpoint.called)

    def test_default_unlimited_and_summary_counts(self):
        with patch.object(sys, 'argv', ['pdd_orders.py']):
            self.assertEqual(pdd.parse_args().max_rounds, 0)
        orders = [pdd.Order('a', '2026-09-03'), pdd.Order('b'), pdd.Order('c', '2025-12-31')]
        summary = pdd.collection_summary(orders, date(2026, 1, 1), date(2026, 9, 23))
        self.assertEqual(summary, {'total': 3, 'in_range': 1, 'unknown': 1, 'outside': 1,
                                  'oldest': '2025-12-31', 'newest': '2026-09-03'})

    def test_paid_amount_is_not_goods_price_or_due(self):
        text = '订单号：260101-123456789\n待收货\n商品：测试商品\n商品金额 ￥100\n待付金额 ￥20\n实付金额：￥80.50\n下单时间：2026-01-01 12:00:00'
        order = pdd.parse_order_card({'text': text, 'lines': text.splitlines()})
        self.assertEqual(order.amount, '80.50')
        self.assertEqual(order.status, '待收货')
        self.assertEqual(order.product, '测试商品')
        self.assertEqual(pdd.paid_amount_from_text('商品金额 ￥100 待付金额 ￥20'), '')
        self.assertEqual(pdd.paid_amount_from_text('实付 ￥0.00'), '0.00')
        self.assertEqual(pdd.paid_amount_from_json({'pay_amount': 1234}), '')
        self.assertEqual(pdd.paid_amount_from_json({'paid_amount_fen': 1234}), '12.34')

    def test_detail_identity_and_multiple_packages(self):
        c = pdd.Capture()
        oid = '260101-123456789'
        c.merge(pdd.Order(oid, product='商品', status='待收货', amount='12.00'))
        self.assertFalse(pdd.merge_detail_text(c, oid, '订单号：260101-999999999\n实付 ￥999'))
        self.assertTrue(pdd.merge_detail_text(c, oid, '运单号：SF123456789', bound_logistics=True))
        self.assertTrue(pdd.merge_detail_text(c, oid, '运单号：YT987654321', bound_logistics=True))
        self.assertEqual(c.orders[oid].tracking_number, 'SF123456789；YT987654321')
        self.assertEqual(c.orders[oid].amount, '12.00')
        self.assertEqual(c.orders[oid].status, '待收货')

    def test_only_verified_detail_route_is_used(self):
        oid = '260101-123456789'
        url = f'https://mobile.yangkeduo.com/example.html?order_sn={oid}&from=orders'
        t = pdd.detail_template(url, f'订单号：{oid}', [oid])
        self.assertIsNotNone(t)
        self.assertIn('order_sn=next-order', pdd.detail_url(t, 'next-order'))
        self.assertIsNone(pdd.detail_template(url, '订单号：260101-000000000', [oid]))
        self.assertIsNone(pdd.detail_template(url.replace('mobile.yangkeduo.com', 'evil.example'), f'订单号：{oid}', [oid]))

    def test_response_shipping_does_not_cross_orders(self):
        values = pdd.shipping_values({'order_sn': 'a', 'packages': [
            {'tracking_no': 'SF12345678'}, {'order_sn': 'b', 'tracking_no': 'OTHER'}]}, 'a', {'tracking_no'})
        self.assertEqual(values, ['SF12345678'])

    def test_scrolling_long_page_is_not_stalled(self):
        self.assertEqual(pdd.stalled_rounds(4, 0, {'before': 100, 'after': 700}, False), 0)
        self.assertEqual(pdd.stalled_rounds(4, 0, {'before': 700, 'after': 700}, False), 5)
        self.assertEqual(pdd.stalled_rounds(4, 1, {'before': 700, 'after': 700}, False), 0)
        self.assertEqual(pdd.stalled_rounds(4, 0, {'before': 700, 'after': 700}, True), 5)

    def test_dates_do_not_use_shipping_date(self):
        order = pdd.parse_order_card({'text': '订单号：260101-123456789 预计送达 2026-01-04'})
        self.assertEqual(order.order_time, '')
        self.assertFalse(pdd.in_range(order, date(2026, 1, 1), date(2026, 1, 5)))
        self.assertIsNone(pdd.parse_order_card({'text': '快递单号：12345678901234'}))

    def test_boundaries_and_merge(self):
        capture = pdd.Capture()
        for order in pdd.response_orders({'data': [{'order_sn': '260101-123456789', 'order_time': '2026-01-01 00:00:00'}]}):
            capture.merge(order)
        capture.merge(pdd.Order('260101-123456789', product='商品'))
        self.assertEqual(len(capture.orders), 1)
        self.assertTrue(pdd.in_range(next(iter(capture.orders.values())), date(2026, 1, 1), date(2026, 1, 1)))

    def test_export_preserves_ids_and_unknown_dates(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            orders = [pdd.Order('001234567890123456789', '2026-01-01 00:00:00', amount='12.30', product='=1+1'),
                      pdd.Order('unknown', raw_text='送达时间 2026-01-02'),
                      pdd.Order('old', '2025-12-31')]
            path = pdd.write_excel(orders, Path(directory), date(2026, 1, 1), date(2026, 1, 1), '测试', 'test')
            book = load_workbook(path)
            self.assertEqual(book['订单'].max_row, 2)
            self.assertEqual(book['日期待核对'].max_row, 2)
            self.assertEqual(book['订单']['A2'].value, orders[0].order_id)
            self.assertEqual(book['订单']['E2'].data_type, 's')
            self.assertEqual(book['订单']['D2'].value, 12.3)
            self.assertEqual(book['订单']['B2'].value.year, 2026)
            book.close()


if __name__ == '__main__':
    unittest.main()
