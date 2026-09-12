from io import BytesIO
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from bit import bit_prohibited_listing_sync as sync
import bit.bit_interface as workbench


def _client():
    workbench.app.config.update(TESTING=True, SECRET_KEY="test-secret")
    client = workbench.app.test_client()
    with client.session_transaction() as flask_session:
        flask_session["workbench_user"] = {
            "id": 1,
            "username": "tester",
            "display_name": "测试用户",
        }
    return client


@pytest.fixture(autouse=True)
def reset_sync_state():
    with sync._state_guard:
        sync._sync_state.update(
            running=False,
            task_id="",
            status="idle",
            message="等待同步禁限售列表",
            total_stores=0,
            processed_stores=0,
            current_store="",
            active_stores=[],
            scanned_count=0,
            reason_matched_count=0,
            prohibited_count=0,
            detail_failed_count=0,
            failed_count=0,
            results=[],
            logs=[],
        )
    yield


def test_official_api_sync_keeps_only_current_forbidden_items(monkeypatch):
    class FakeClient:
        def request(self, _method, path, params=None):
            if path == "/marketplace/users/root-1":
                return {"marketplaces": [{"site_id": "MLM", "user_id": "seller-mx"}]}
            assert path == "/marketplace/moderations/infractions/seller-mx"
            if params["offset"] == 0:
                return {
                    "infractions": [
                        {
                            "id": "inf-1",
                            "related_item_id": "MLM1",
                            "reason": "The product is prohibited.",
                            "date_created": "2026-08-29T10:20:30Z",
                        },
                        {
                            "id": "inf-2",
                            "related_item_id": "MLM2",
                            "reason": "The product's brand is not generic.",
                        },
                    ],
                    "paging": {"total": 3},
                }
            return {
                "infractions": [{
                    "id": "inf-3",
                    "related_item_id": "MLM3",
                    "reason": "The product is prohibited.",
                }],
                "paging": {"total": 3},
            }

        def get_marketplace_item(self, item_id, attributes=None):
            assert attributes
            if item_id == "MLM1":
                return {
                    "id": item_id,
                    "site_id": "MLM",
                    "seller_id": "seller-mx",
                    "title": "当前禁售商品",
                    "status": "under_review",
                    "sub_status": ["forbidden"],
                    "user_product_id": "UP-1",
                }
            return {
                "id": item_id,
                "site_id": "MLM",
                "seller_id": "seller-mx",
                "title": "已经恢复的商品",
                "status": "active",
                "sub_status": [],
            }

    captured = {}
    monkeypatch.setattr(sync, "PROHIBITED_PAGE_SIZE", 2)
    monkeypatch.setattr(sync, "_client_and_token", lambda record: (FakeClient(), record))
    monkeypatch.setattr(sync, "get_prohibited_sync_context", lambda _token_id: {})
    monkeypatch.setattr(
        sync,
        "replace_prohibited_snapshot",
        lambda token, rows, **kwargs: captured.update(
            token=dict(token), rows=list(rows), kwargs=kwargs
        ) or {"total": len(captured["rows"])},
    )
    result = sync._sync_store({
        "id": 7,
        "display_name": "顺风顺水（fti）",
        "meli_user_id": "root-1",
        "access_token": "secret",
        "site_settings": [{
            "site_id": "MLM",
            "salesperson": "业务员甲",
            "group_name": "墨西哥组",
        }],
    })

    assert result["scanned"] == 3
    assert result["reason_matched"] == 2
    assert result["prohibited"] == 1
    assert captured["rows"][0]["item_id"] == "MLM1"
    assert captured["rows"][0]["global_item_id"] == "UP-1"
    assert captured["rows"][0]["salesperson"] == "业务员甲"
    assert captured["rows"][0]["infraction_date"] == "2026-08-29 10:20:30"


def test_prohibited_listing_ui_and_list_route():
    client = _client()
    response = client.get("/")
    assert response.status_code == 200
    assert b'data-tab="prohibited-listings"' in response.data
    assert b'id="tab-prohibited-listings"' in response.data
    assert "禁限售列表".encode("utf-8") in response.data
    assert b"The product is prohibited." in response.data
    assert b'id="prohibited-risk-type-filter"' in response.data
    assert b'id="prohibited-reply-count"' in response.data
    assert b'id="prohibited-auto-sync-enabled"' in response.data
    assert "最新更新时间".encode("utf-8") in response.data
    assert b'scope: "prohibited_listings"' in response.data
    assert "待回复权利人".encode("utf-8") in response.data
    assert b"startProhibitedListingSync" in response.data

    listing_data = {
        "rows": [{"item_id": "MLM1", "infraction_reason": sync.PROHIBITED_REASON}],
        "groups": [],
        "stores": [],
        "salespersons": [],
        "summary": {
            "current_count": 2,
            "prohibited_count": 1,
            "rights_holder_reply_count": 1,
        },
        "total": 1,
        "page": 1,
        "pages": 1,
        "page_size": 100,
    }
    with patch.object(
        workbench.bit_db_api,
        "list_mercado_prohibited_listings",
        return_value=listing_data,
    ) as listing:
        response = client.get(
            "/api/prohibited-listings?token_id=7&site_id=MLM&salesperson="
            "%E4%B8%9A%E5%8A%A1%E5%91%98%E7%94%B2&risk_type=rights_holder_reply"
            "&search=MLM1&page=1&page_size=100"
        )
    assert response.status_code == 200
    assert response.get_json()["data"]["rows"][0]["item_id"] == "MLM1"
    assert listing.call_args.kwargs["token_id"] == 7
    assert listing.call_args.kwargs["site_id"] == "MLM"
    assert listing.call_args.kwargs["salesperson"] == "业务员甲"
    assert listing.call_args.kwargs["risk_type"] == "rights_holder_reply"


def test_prohibited_listing_export_keeps_filters_and_exports_all_pages():
    received = []

    def fake_list(**kwargs):
        received.append(kwargs)
        page = int(kwargs["page"])
        return {
            "rows": [{
                "risk_type": "prohibited" if page == 1 else "rights_holder_reply",
                "salesperson": "业务员甲",
                "group_name": "墨西哥组",
                "store_name": "测试店铺",
                "site_id": "MLM",
                "risk_count": 3,
                "prohibited_count": 2,
                "rights_holder_reply_count": 1,
                "item_id": f"MLM{page}",
                "global_item_id": "UP-1",
                "title": '=HYPERLINK("bad")' if page == 1 else "普通商品",
                "status": "under_review",
                "sub_status": ["forbidden"],
                "infraction_reason": "The product is prohibited.",
                "remedy": "请检查商品资质",
                "rights_holder": "权利人甲" if page == 2 else "",
                "due_at": "2026-09-20 00:00:00" if page == 2 else None,
                "infraction_date": "2026-09-08 10:20:30",
                "last_checked_at": "2026-09-08 10:30:30",
                "permalink": "https://example.test/item/MLM1",
                "thumbnail_url": "https://example.test/image/MLM1.jpg",
            }],
            "page": page,
            "pages": 2,
        }

    client = _client()
    with patch.object(workbench.bit_db_api, "list_mercado_prohibited_listings", side_effect=fake_list):
        response = client.get(
            "/api/prohibited-listings/export",
            query_string={
                "token_id": 7,
                "site_id": "MLM",
                "salesperson": "业务员甲",
                "risk_type": "prohibited",
                "search": "MLM",
            },
        )

    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert len(received) == 2
    assert received[0]["token_id"] == 7
    assert received[0]["site_id"] == "MLM"
    assert received[0]["salesperson"] == "业务员甲"
    assert received[0]["risk_type"] == "prohibited"
    assert received[0]["page_size"] == 500
    assert received[1]["page"] == 2
    workbook = load_workbook(BytesIO(response.data))
    sheet = workbook["禁限售明细"]
    assert sheet.max_row == 3
    assert sheet["A2"].value == "当前禁售"
    assert sheet["A3"].value == "待回复权利人"
    assert sheet["K2"].value.startswith("'=HYPERLINK")
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:U3"


def test_prohibited_listing_ui_exposes_excel_export_button():
    source = workbench.resolve_template_dir() / "index.html"
    html = source.read_text(encoding="utf-8")
    assert 'id="prohibited-export-button"' in html
    assert "exportProhibitedListings" in html
    assert "/api/prohibited-listings/export" in html


def test_manual_prohibited_sync_route_can_target_one_store():
    client = _client()
    with patch.object(
        workbench.bit_db_api,
        "start_prohibited_listing_sync",
        return_value={"started": True, "state": {"running": True, "task_id": "p-1"}},
    ) as start:
        response = client.post(
            "/api/prohibited-listings/sync/start",
            json={"token_ids": [7]},
        )
    assert response.status_code == 202
    assert response.get_json()["data"]["running"] is True
    start.assert_called_once_with([7])


def test_due_prohibited_sync_respects_disabled_full_refresh_switch(monkeypatch):
    monkeypatch.setattr(
        sync,
        "get_overview_sync_settings",
        lambda _scope: {"enabled": False},
    )
    monkeypatch.setattr(
        sync,
        "list_due_prohibited_token_ids",
        lambda **_kwargs: pytest.fail("disabled scheduler must not query due stores"),
    )

    result = sync.start_due_prohibited_listing_sync()

    assert result["disabled"] is True
    assert result["started"] is False


def test_all_store_prohibited_sync_ignores_selected_ids():
    client = _client()
    with patch.object(
        workbench.bit_db_api,
        "start_prohibited_listing_sync",
        return_value={"started": True, "state": {"running": True}},
    ) as start:
        response = client.post(
            "/api/prohibited-listings/sync/start",
            json={"sync_all": True, "token_ids": [7]},
        )
    assert response.status_code == 202
    start.assert_called_once_with([])


def test_salesperson_sync_resolves_authorized_store_ids():
    client = _client()
    tokens = {
        "rows": [
            {"id": 2, "site_settings": [{"site_id": "MLM", "salesperson": "张泽文"}]},
            {"id": 3, "site_settings": [{"site_id": "MLB", "salesperson": "其他人"}]},
            {"id": 4, "site_settings": [{"site_id": "MCO", "salesperson": "张泽文"}]},
        ]
    }
    with (
        patch.object(workbench.bit_db_api, "list_mercado_store_tokens", return_value=tokens),
        patch.object(
            workbench.bit_db_api,
            "start_prohibited_listing_sync",
            return_value={"started": True, "state": {"running": True}},
        ) as start,
    ):
        response = client.post(
            "/api/prohibited-listings/sync/start",
            json={"salesperson": "张泽文"},
        )
    assert response.status_code == 202
    start.assert_called_once_with([2, 4])
