"""离线浏览器流程测试：路由全部拦截，不连接真实拼多多账号。"""
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch
from urllib.parse import parse_qs, urlsplit

from test_pdd_orders_export import pdd


class DetailBrowserTests(unittest.TestCase):
    def test_launch_and_reuse_dynamic_edge_port(self):
        from playwright.sync_api import sync_playwright
        original = pdd.edge_base_args
        # Edge 子进程在正常关闭后可能短暂占用资料文件，不影响连接复用断言。
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            profile = Path(directory)
            with patch.object(pdd, 'edge_base_args', side_effect=lambda path: original(path) + ['--headless=new']), \
                 patch.object(pdd, 'ORDERS_URL', 'about:blank'):
                process, endpoint = pdd.launch_edge_for_collection(pdd.find_edge_executable(), profile)
            with sync_playwright() as runtime:
                browser = runtime.chromium.connect_over_cdp(endpoint)
                try:
                    with patch.object(pdd.subprocess, 'Popen') as launch:
                        second, second_endpoint = pdd.launch_edge_for_collection(pdd.find_edge_executable(), profile)
                    self.assertIsNone(second)
                    self.assertEqual(second_endpoint, endpoint)
                    launch.assert_not_called()
                finally:
                    session = browser.new_browser_cdp_session()
                    session.send('Browser.close')
                    process.wait(timeout=15)

    def test_two_orders_and_logistics_export(self):
        from playwright.sync_api import sync_playwright
        from openpyxl import load_workbook

        ids = ['260101-123456789', '260102-987654321']
        capture = pdd.Capture()
        for oid in ids:
            capture.merge(pdd.Order(oid))

        def respond(route):
            parts = urlsplit(route.request.url)
            oid = parse_qs(parts.query)['order_sn'][0]
            if parts.path == '/fixture-detail.html':
                html = (f'<div>订单号：{oid}</div><div>下单时间：2026-01-02 12:00:00</div>'
                        '<div>待收货</div><div>商品：测试商品</div><div>商品金额 ￥100</div>'
                        '<div>实付金额：￥80.50</div>'
                        f'<a href="/fixture-logistics.html?order_sn={oid}">查看物流</a>')
            else:
                html = f'<div>快递单号：SF{oid.replace("-", "")}</div><div>顺丰</div>'
            route.fulfill(status=200, content_type='text/html; charset=utf-8', body=html)

        with sync_playwright() as runtime:
            browser = runtime.chromium.launch(executable_path=str(pdd.find_edge_executable()), headless=True)
            try:
                page = browser.new_page()
                page.set_content('<main>' + ''.join(
                    f'<section><div><span>订单号：</span><span>{oid}</span></div><div>商品：测试商品</div></section>'
                    for oid in ids) + '</main>')
                cards = page.evaluate(pdd.EXTRACT_CARDS_JS)
                self.assertEqual({pdd.parse_order_card(card).order_id for card in cards}, set(ids))
                page.route('**/*', respond)
                page.goto(f'https://mobile.yangkeduo.com/fixture-detail.html?order_sn={ids[0]}')
                with patch('builtins.input', return_value=''):
                    pdd.enrich_orders(page, date(2026, 1, 1), date(2026, 1, 3), capture, lambda: None, 0)
                for oid in ids:
                    order = capture.orders[oid]
                    self.assertEqual(order.amount, '80.50')
                    self.assertEqual(order.status, '待收货')
                    self.assertEqual(order.product, '测试商品')
                    self.assertEqual(order.tracking_number, 'SF' + oid.replace('-', ''))
                    self.assertEqual(pdd.missing_fields(order), '')
                with tempfile.TemporaryDirectory() as directory:
                    path = pdd.write_excel(list(capture.orders.values()), Path(directory), date(2026, 1, 1), date(2026, 1, 3), '测试', 'v2')
                    book = load_workbook(path)
                    self.assertEqual(book['订单'].max_row, 3)
                    self.assertEqual(book['订单']['D1'].value, '实付金额（元）')
                    self.assertEqual(book['订单']['F2'].value, 'SF' + ids[0].replace('-', ''))
                    book.close()
            finally:
                browser.close()


if __name__ == '__main__':
    unittest.main()
