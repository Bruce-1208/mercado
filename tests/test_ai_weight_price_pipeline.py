"""Offline pipeline scenarios; never use real suppliers or write to ERP."""
import io
import threading

import pytest
from openpyxl import load_workbook

from erp.ai_weight_price.config import validate as validate_config, selection_key
from erp.ai_weight_price.models import Models
from erp.ai_weight_price.reports import execution_xlsx
from erp.ai_weight_price.service import Service, ItemBlocked
from erp.ai_weight_price.store import Store

def validate(value):
    """Legacy consultation coverage; image-first behavior has dedicated tests."""
    return validate_config({"workflow_mode": "legacy_consult", **value})


class PageModel:
    def match(self, task, candidate):
        return {**candidate, "selected_sku": candidate["skus"][0], "confidence": .98}, [{"confirmed": True}]

    def supplier_info(self, task, text, source="1688页面"):
        return {"weight_g": "450" if "包装450g" in text else None,
                "cost_price": "22" if "最终22元" in text else None}

    def weight(self, text):
        return "450" if "包装450g" in text else None


class PipelineBrowser:
    def __init__(self, config, stop, log):
        self.config, self.stop = config, stop
        self.operations, self.messages = [], []
        self.description, self.price = "包装450g", "22"
        self.fail_write = False

    def __enter__(self): return self
    def __exit__(self, *args): pass
    def confirm_login(self): pass
    def release(self, *args): pass

    def collect(self, store, on_task):
        for index in range(12):
            key = str(index + 1)
            self.operations.append(("collect", key))
            store.add({"erp_goods_id": key, "title": "测试商品" + key, "erp_sku": "蓝色一件",
                       "main_image_url": "https://img.example/" + key + ".jpg", "reference_weight_g": "450",
                       "erp_edit_url": "https://meli.zying.net/#/product/" + key})
            store.include_in_scope(self.config["run_scope"], key, 1)
            on_task(key)

    def candidates(self, task):
        key = task["erp_goods_id"]
        self.operations.append(("image", key))
        yield {"url": "https://detail.1688.com/offer/" + key + ".html", "merchant_id": "m" + key,
               "description": self.description, "skus": [{"id": "sku", "label": "蓝色一件", "price": self.price}]}

    def prepare_chat(self, task): return object(), "https://air.1688.com/chat/" + task["merchant_id"], []
    def send(self, page, task, message): self.messages.append(message)
    def replies(self, task): return [{"id": "reply", "text": "包装450g，最终22元", "at": task["sent_at"] + 1}]

    def write(self, task, before_save):
        key = task["erp_goods_id"]
        before_save({"net_income_usd": "9.5", "weight_g": "430"})
        self.operations.append(("save", key))
        if self.fail_write:
            raise ValueError("模拟保存失败")
        return {"net_income_usd": task["net_income_usd"], "weight_g": task["weight_g"]}


def setup(tmp_path, monkeypatch):
    config = validate({"writeback_enabled": True, "usd_cny_rate": "7.2"})
    config.update(run_selection={"category": "", "start_page": 1, "end_page": 1}, max_items=10, run_id="offline-ten")
    config["run_scope"] = selection_key(config["run_selection"], config)
    browser = PipelineBrowser(config, threading.Event(), lambda *args: None)
    service = Service(tmp_path, browser_factory=lambda *args: browser, models_factory=lambda *args: PageModel())
    monkeypatch.setattr(service, "exchange_rate", lambda config: {"cny_per_usd": "7.2", "date": "2026-09-07", "source": "manual"})
    service.store.set_state("run", {"run_id": config["run_id"], "processed_items": 0})
    return service, browser, config


def add_task(service):
    service.store.add({"erp_goods_id": "1", "title": "测试商品", "erp_sku": "蓝色一件", "reference_weight_g": "450",
                       "erp_edit_url": "https://meli.zying.net/#/product/1"})
    return service.store.get("1")


def test_ten_products_strict_order_persistent_database_logs_and_excel(tmp_path, monkeypatch):
    service, browser, config = setup(tmp_path, monkeypatch)
    lock = service.lock()
    assert lock.acquire()
    service.run(config, "pipeline", None, lock)
    assert browser.operations == [(step, str(index)) for index in range(1, 11) for step in ("collect", "image", "save")]
    assert not browser.messages
    persisted = Store(tmp_path)
    assert persisted.counts()["success"] == 10
    assert persisted.list()["total"] == 10  # eleventh product was never collected
    run, rows = persisted.run_report("offline-ten")
    assert run["outcome"] == "completed" and len(rows) == 10
    assert all(row["erp_before"] == {"weight_g": "430", "net_income_usd": "9.5"} for row in rows)
    assert all(row["erp_after"] == {"weight_g": "450", "net_income_usd": "4"} and row["write_verified"] for row in rows)
    assert sum("回填前：" in log["message"] for log in persisted.logs(limit=1000)) == 10
    assert sum("保存成功并回读确认" in log["message"] for log in persisted.logs(limit=1000)) == 10
    assert [log["task_id"] for log in persisted.logs(limit=1000)
            if log["message"].startswith("开始逐件核对第 ")] == list(map(str, range(1, 11)))
    sheet = load_workbook(io.BytesIO(execution_xlsx(rows, run))).active
    assert sheet.max_row == 15 and sheet.freeze_panes == "C6"
    assert sheet["I6"].value == 430 and sheet["J6"].value == 9.5
    assert sheet["M6"].value == 450 and sheet["N6"].value == 4


@pytest.mark.parametrize("accept_even", [False, True])
def test_unmatched_products_continue_and_count_toward_ten_with_excel(tmp_path, monkeypatch, accept_even):
    service, browser, config = setup(tmp_path, monkeypatch)
    class SelectiveModel(PageModel):
        def match(self, task, candidate):
            if accept_even and int(task["erp_goods_id"]) % 2 == 0:
                return super().match(task, candidate)
            return None, [{"confidence": .6, "reason": "颜色和包装数量不一致"}]
    service.models_factory = lambda *args: SelectiveModel()
    lock = service.lock(); assert lock.acquire()
    service.run(config, "pipeline", None, lock)
    store = Store(tmp_path)
    counts = store.counts()
    assert counts["skipped"] == (5 if accept_even else 10)
    assert counts["success"] == (5 if accept_even else 0)
    assert store.list()["total"] == 10
    assert [key for step, key in browser.operations if step == "image"] == list(map(str, range(1, 11)))
    assert not browser.messages
    assert store.state("pipeline_current") is None
    run, rows = store.run_report(config["run_id"])
    assert run["processed_items"] == 10 and run["skipped_items"] == counts["skipped"]
    assert run["outcome"] == "completed" and len(rows) == 10
    assert sum("未完全匹配，已跳过" in log["message"] for log in store.logs(limit=1000)) == counts["skipped"]
    assert rows[0]["execution_result"] == "已跳过（未完全匹配）"
    assert rows[0]["match_evidence"][0]["reviews"][0]["reason"] == "颜色和包装数量不一致"
    sheet = load_workbook(io.BytesIO(execution_xlsx(rows, run))).active
    assert sheet["C6"].value == "已跳过（未完全匹配）"
    assert "未通过" in sheet["D6"].value
    assert sheet["I6"].value is None and sheet["M6"].value is None


def test_skipped_products_do_not_block_resume_and_can_be_retried(tmp_path, monkeypatch):
    service, browser, config = setup(tmp_path, monkeypatch)
    add_task(service)
    service.store.skip("1", "无完全匹配")
    service.store.set_state("pipeline_current", {"scope": config["run_scope"], "task_id": "1"})
    service.complete_one("1", browser, PageModel(), config)
    assert not browser.operations and service.store.state("pipeline_current") is None
    service.retry("1")
    task = service.store.get("1")
    assert task["status"] == "pending" and task["skip_reason"] == ""
    assert task["retry_history"][-1]["reason"] == "无完全匹配"


def test_empty_search_is_skipped_without_changing_erp(tmp_path, monkeypatch):
    from erp.ai_weight_price.browser import NoExactMatch
    service, browser, config = setup(tmp_path, monkeypatch)
    add_task(service)
    monkeypatch.setattr(browser, "candidates", lambda task: (_ for _ in ()).throw(NoExactMatch("1688未找到相关商品")))
    service.complete_one("1", browser, PageModel(), config)
    assert service.store.get("1")["status"] == "skipped"
    assert not browser.operations and not browser.messages


@pytest.mark.parametrize("page_text,page_price,missing", [("", "22", "重量"), ("包装450g", None, "供货成本"), ("", None, "重量")])
def test_only_missing_fields_use_chat_and_preserve_page_facts(tmp_path, monkeypatch, page_text, page_price, missing):
    service, browser, config = setup(tmp_path, monkeypatch)
    browser.description, browser.price = page_text, page_price
    service.process(add_task(service), browser, PageModel(), config)
    waiting = service.store.get("1")
    assert waiting["status"] == "waiting_merchant_reply" and missing in browser.messages[0]
    assert not any(step == "save" for step, key in browser.operations)
    service.poll(waiting, browser, PageModel(), config, waiting["next_poll_at"])
    saved = service.store.get("1")
    assert saved["status"] == "success" and saved["net_income_usd"] == "4"
    if page_text:
        assert saved["info_sources"]["weight_g"] == "1688页面"


def test_failure_isolated_to_item_and_next_collection_continues(tmp_path, monkeypatch):
    service, browser, config = setup(tmp_path, monkeypatch)
    browser.fail_write = True
    lock = service.lock(); assert lock.acquire()
    service.run(config, "pipeline", None, lock)
    assert service.store.list()["total"] == 10
    assert service.store.state("run")["outcome"] == "completed"
    assert service.store.state("run")["processed_items"] == 10
    assert service.store.get("1")["status"] == "skipped"
    assert service.store.get("1")["erp_before"]["weight_g"] == "430"
    assert service.store.get("1")["erp_after"] is None
    assert not service.store.get("1")["write_verified"]
    assert [key for step, key in browser.operations if step == "collect"] == list(map(str, range(1, 11)))


def test_waiting_current_item_never_processes_next_item(tmp_path, monkeypatch):
    service, browser, config = setup(tmp_path, monkeypatch)
    browser.description = ""
    service.process(add_task(service), browser, PageModel(), config)
    service.store.include_in_scope(config["run_scope"], "1", 1)
    service.store.add({"erp_goods_id": "2", "title": "下一件"})
    service.store.include_in_scope(config["run_scope"], "2", 1)
    before = list(browser.operations)
    service.tick(browser, PageModel(), config)
    assert browser.operations == before
    assert service.store.get("2")["stage"] == "collected"


def test_failed_single_retry_keeps_original_scope_barrier(tmp_path, monkeypatch):
    service, browser, config = setup(tmp_path, monkeypatch)
    add_task(service)
    service.store.exception("1", "需人工核对")
    service.store.set_state("pipeline_current", {"scope": config["run_scope"], "task_id": "1"})
    with pytest.raises(ItemBlocked):
        service.complete_one("1", browser, PageModel(), validate({}))
    assert service.store.state("pipeline_current")["scope"] == config["run_scope"]


def test_page_complete_does_not_require_chat_preflight(tmp_path, monkeypatch):
    service, browser, config = setup(tmp_path, monkeypatch)
    config = validate({"api_base_url": "http://localhost:11434/v1", "selectors": {
        field: ".test" for field in ("supplier_title", "supplier_image", "supplier_merchant", "sku_rows", "sku_label")}})
    service.preflight(config, "pipeline")


def test_model_requires_explicit_quoted_evidence(tmp_path, monkeypatch):
    model = Models(validate({}), lambda *args: None)
    monkeypatch.setattr(model, "call", lambda *args, **kwargs: {"weight_g": "450", "weight_evidence": "想象的包装重量",
                                                              "cost_price": "22", "cost_evidence": "最终22元"})
    result = model.supplier_info({"erp_sku": "蓝色一件"}, "最终22元")
    assert result["weight_g"] is None and result["cost_price"] == "22"


def test_excel_preserves_ids_and_never_executes_product_text():
    sheet = load_workbook(io.BytesIO(execution_xlsx([{"erp_goods_id": "001234567890123456", "title": '=HYPERLINK("bad")',
        "status": "pending", "execution_result": "启动受阻", "execution_reason": "字段缺失"}]))).active
    assert sheet["A6"].value == "001234567890123456"
    assert sheet["B6"].data_type == "s"
    assert sheet["I6"].value is None and sheet["M6"].value is None
