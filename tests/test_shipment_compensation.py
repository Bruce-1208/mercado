import io
import json
from decimal import Decimal
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from bit import bit_mysql as db, weight_dimensions_records as records
from bit.shipment_compensation import apply_dimension_costs, dimension_adjustment


def payload(entries, cost=10):
    return {"currency_id": "USD", "senders": [{"cost": cost, "compensations": entries}]}


@pytest.mark.parametrize("entries, expected", [
    ([], "0"),
    ([{"reason": "other", "amount": 5}], "0"),
    ([{"reason": "incorrect_dimensions", "amount": 0}], "0"),
    ([{"reason": "incorrect_dimensions", "amount": "NaN"}], "0"),
    ([{"reason": "incorrect_dimensions", "amount": 2.5}], "2.5"),
    ([{"reason": "incorrect_dimensions", "amount": -2.5}], "-2.5"),
    ([{"reason": "incorrect_dimensions", "amount": 2},
      {"reason": "incorrect_dimensions", "amount": -2}], "0"),
    ([{"reason": "incorrect_dimensions", "amount": 2},
      {"reason": "other", "amount": 9}], "2"),
])
def test_only_nonzero_measurement_adjustments_qualify(entries, expected):
    data = payload(entries)
    assert dimension_adjustment(data) == Decimal(expected)
    row = {"freight_difference": "-100 USD"}
    apply_dimension_costs(row, data)
    assert row["_dimension_compensation_confirmed"] == bool(Decimal(expected))
    if Decimal(expected):
        assert row["freight_difference_usd"] == expected
        assert row["declared_freight"] == "10 USD"
        assert Decimal(row["actual_freight"].split()[0]) == 10 + Decimal(expected)


def test_background_candidates_do_not_require_a_local_freight_quote():
    data = payload([{"reason": "incorrect_dimensions", "amount": 2}])
    with patch.object(db, "list_orders", return_value={
        "rows": [{"id": "child", "order_number": "pack", "store_id": 1,
                  "platform_shipping_id": "shipment", "actual_freight_usd": None,
                  "quoted_freight_usd": None}], "total": 1, "page_size": 200,
    }) as orders, patch.object(db.pymysql, "connect") as connect:
        cursor = connect.return_value.cursor.return_value.__enter__.return_value
        cursor.fetchall.side_effect = [[], [{"order_id": "child", "raw_json": "{}",
                                            "payload_json": json.dumps(data)}]]
        rows = db.list_weight_dimensions_changed_orders()
    assert orders.call_args.kwargs["freight_variance"] == "dimensions_changed"
    assert len(rows) == 1
    assert rows[0]["freight_difference"] == "2 USD"
    assert rows[0]["_dimension_compensation_confirmed"] is True


def test_saved_records_recompute_stale_freight_and_use_same_filter_for_count_and_rows():
    data = payload([{"reason": "incorrect_dimensions", "amount": -2}])
    with patch.object(db.pymysql, "connect") as connect, patch.object(db, "_ensure_weight_dimensions_record_table"):
        cursor = connect.return_value.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = {"total": 1}
        cursor.fetchall.return_value = [{
            "order_number": "1", "record_json": json.dumps({"freight_difference": "99 USD"}),
            "execution_logs_json": "[]", "compensation_payload_json": json.dumps(data),
        }]
        result = db.list_weight_dimensions_records(filters={"freight_min": -3, "freight_max": -1}, page=1)
    count_sql, row_sql = [call.args[0] for call in cursor.execute.call_args_list]
    count_where = count_sql.split("WHERE ", 1)[1]
    assert row_sql.endswith(count_where + "ORDER BY `wdr_freight_changed_at` DESC, `id` DESC LIMIT %s OFFSET %s")
    assert "incorrect_dimensions" in count_where
    assert "$.freight_difference_usd" not in count_where
    assert result["records"][0]["freight_difference"] == "-2 USD"


def test_upload_pagination_and_export_exclude_unconfirmed_adjustments():
    task = {"task_id": "t", "owner": "u", "source": "upload", "status": "ready", "records": [
        {"order_number": "quote-only"},
        {"order_number": "weight-only", "actual_weight_g": "400", "actual_dimensions_cm": "",
         "_dimension_compensation_confirmed": True},
        {"order_number": "zero", "_dimension_compensation_confirmed": False},
    ]}
    with patch.object(records, "_tasks", {"t": task}):
        result = records.task_status("t", "u", page=1)
        exported = records.export_xlsx("t", "u")
    assert result["record_total"] == 1
    assert result["records"][0]["order_number"] == "weight-only"
    book = load_workbook(io.BytesIO(exported))
    assert book.active.max_row == 2
    book.close()


def test_measurement_refresh_preserves_official_dimension_costs():
    from unittest.mock import Mock
    data = payload([{"reason": "incorrect_dimensions", "amount": 3}])
    row = {}
    apply_dimension_costs(row, data)
    client = Mock()
    client.request.return_value = {"costs": {"currency_id": "USD", "validated": {"sender": {"cost": 99}}}}
    client.get_shipment.return_value = {"dimensions": {}}
    records._read_measurements(row, client, "shipment")
    assert row["actual_freight"] == "13 USD"
    assert row["freight_difference"] == "3 USD"
